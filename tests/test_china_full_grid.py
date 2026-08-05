from __future__ import annotations

import hashlib
import importlib.util
import json
import math
import subprocess
import sys
from pathlib import Path

import geopandas as gpd
import pandas as pd
import pyarrow.parquet as pq
import pytest
from docx import Document
from docx.oxml.ns import qn
from pyproj import Transformer
from shapely import normalize, set_precision
from shapely.affinity import translate
from shapely.geometry import box

MODULE_PATH = Path(__file__).parents[1] / "scripts/data/china_full_grid.py"
SPEC = importlib.util.spec_from_file_location("china_full_grid", MODULE_PATH)
assert SPEC is not None and SPEC.loader is not None
MODULE = importlib.util.module_from_spec(SPEC)
sys.modules[SPEC.name] = MODULE
SPEC.loader.exec_module(MODULE)

DOCX_MODULE_PATH = Path(__file__).parents[1] / "scripts/docs/build_china_full_grid_package_docx.py"
PACKAGE_README_PATH = (
    Path(__file__).parents[1] / "docs/data/china_full_1280m_grid_package_readme_20260805.md"
)
BUILDER_MODULE_PATH = Path(__file__).parents[1] / "scripts/data/build_china_full_grid.py"


def _load_module(path: Path, name: str):
    assert path.exists(), f"expected Task 4 module at {path}"
    spec = importlib.util.spec_from_file_location(name, path)
    assert spec is not None and spec.loader is not None
    module = importlib.util.module_from_spec(spec)
    sys.modules[name] = module
    spec.loader.exec_module(module)
    return module


def _macro() -> dict[str, int | str]:
    return {
        "grid_epsg": 32650,
        "grid_id": "utm50n",
        "macro_col": 34,
        "macro_row": 344,
    }


def _boundary():
    return box(116.28, 39.83, 116.31, 39.86)


def synthetic_parent_records(count: int) -> list[dict[str, object]]:
    """Return canonical records spanning two macro cells for writer tests."""
    spec = MODULE.GridSpec(boundary_version="test")
    first_macro = _macro()
    second_macro = {**first_macro, "macro_col": first_macro["macro_col"] + 1}
    records: list[dict[str, object]] = []
    for index in range(count):
        macro = first_macro if index < count // 2 else second_macro
        grid_col = int(macro["macro_col"]) * spec.macro_side_patches + index % 10
        grid_row = int(macro["macro_row"]) * spec.macro_side_patches + index // 10
        records.append(
            MODULE.build_patch_record(
                macro,
                grid_col,
                grid_row,
                116.28 + index * 0.001,
                39.83 + index * 0.001,
                spec,
            )
        )
    return records


def _valid_package_audit(
    *, all_count: int = 12, sampled_count: int = 2, unsampled_count: int = 10
) -> dict[str, object]:
    return {
        "schema_version": "china_full_1280m_membership_audit_v1",
        "passed": True,
        "all_count": all_count,
        "sampled_count": sampled_count,
        "unsampled_count": unsampled_count,
        "matched": sampled_count,
        "missing_sampled_count": 0,
        "duplicate_parent_key_count": 0,
        "sampled_unsampled_intersection_count": 0,
        "sampled_flag_mismatch_count": 0,
        "partition_mismatch_count": 0,
        "exact_partition_membership_mismatch_count": 0,
        "child_partition_unknown_parent_key_count": 0,
        "child_partition_metadata_mismatch_count": 0,
        "child_partition_geometry_mismatch_count": 0,
        "footprint_coordinate_mismatch_count": 0,
        "stored_utm_bounds_mismatch_count": 0,
        "stored_wgs84_bounds_mismatch_count": 0,
        "invalid_geometry_count": 0,
        "owner_zone_mismatch_count": 0,
        "same_zone_positive_overlap_count": 0,
        "cross_zone_overlap_violation_count": 0,
        "hash_mismatches": {
            "identity_hash": 0,
            "footprint_hash": 0,
            "sampled_registry_footprint_hash": 0,
        },
    }


def _assert_package_artifacts_absent(output_root: Path) -> None:
    names = {
        "china_full_1280m_macrocell_index.gpkg",
        "national_patch_density_by_utm_zone.png",
        "local_1280m_grid_sampled_unsampled.png",
        "utm_owner_zone_seams.png",
        "README.md",
        "先读我.docx",
        "china_full_1280m_grid_package_manifest.json",
    }
    assert all(not (output_root / name).exists() for name in names)


def test_package_docs_explain_parent_sampled_and_unsampled(tmp_path: Path) -> None:
    docx_module = _load_module(DOCX_MODULE_PATH, "china_full_grid_package_docx")
    assert PACKAGE_README_PATH.exists()
    readme_text = PACKAGE_README_PATH.read_text(encoding="utf-8")
    assert "全国完整父网格" in readme_text
    assert "已采样 patch" in readme_text
    assert "未采样 patch" in readme_text
    prior_sampling_path = tmp_path / "prior_sampling.md"
    prior_sampling_path.write_text(
        "# 先前采样说明\n\n版本：2026-07-27\n\n62,000 个候选位置用于数据获取和质量复核。\n",
        encoding="utf-8",
    )
    docx_path = docx_module.build_package_docx(
        tmp_path / "先读我.docx",
        package_metadata={
            "all_count": 12,
            "sampled_count": 2,
            "unsampled_count": 10,
            "membership_audit": {"passed": True, "matched": 2},
        },
        prior_sampling_source=prior_sampling_path,
    )

    document = Document(docx_path)
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    assert "全国完整父网格" in text
    assert "已采样 patch" in text
    assert "未采样 patch" in text
    assert "EPSG:32643" in text
    assert "X=经度、Y=纬度" in text
    assert "62,000 个候选位置" in text

    title = document.styles["Title"]
    heading = document.styles["Heading 1"]
    normal = document.styles["Normal"]
    assert title.font.size.pt == 16
    assert heading.font.size.pt == 14
    assert normal.font.size.pt == 12
    for style in (title, heading, normal):
        fonts = style._element.rPr.rFonts
        assert fonts.get(qn("w:ascii")) == "Times New Roman"
        assert fonts.get(qn("w:eastAsia")) == "SimSun"
    for paragraph in document.paragraphs:
        for run in paragraph.runs:
            assert str(run.font.color.rgb) == "000000"


def test_package_docs_use_default_prior_sampling_source_in_clean_checkout(tmp_path: Path) -> None:
    docx_module = _load_module(DOCX_MODULE_PATH, "china_full_grid_package_docx_default")

    assert docx_module.DEFAULT_PRIOR_SAMPLING_SOURCE.exists()
    docx_path = docx_module.build_package_docx(
        tmp_path / "先读我.docx",
        package_metadata={
            "all_count": 12,
            "sampled_count": 2,
            "unsampled_count": 10,
            "membership_audit": {"passed": True, "matched": 2},
        },
    )

    text = "\n".join(paragraph.text for paragraph in Document(docx_path).paragraphs)
    assert "62,000 个全国采样候选位置" in text


def test_package_artifacts_build_index_previews_and_manifest_from_synthetic_grid(
    tmp_path: Path,
) -> None:
    builder = _load_module(BUILDER_MODULE_PATH, "build_china_full_grid")
    records = synthetic_parent_records(count=12)
    output_root = tmp_path / "package"
    MODULE.write_zone_records(
        records,
        {str(records[1]["parent_key"]), str(records[7]["parent_key"])},
        output_root,
        batch_size=5,
    )
    audit_path = tmp_path / "membership_audit.json"
    audit_path.write_text(
        json.dumps(_valid_package_audit()),
        encoding="utf-8",
    )
    prior_sampling_path = tmp_path / "prior_sampling.md"
    prior_sampling_path.write_text(
        "# 先前采样说明\n\n62,000 个候选位置用于数据获取和质量复核。\n",
        encoding="utf-8",
    )

    manifest_path = builder.build_package_artifacts(
        output_root=output_root,
        boundary_wgs84=box(113.95, 39.7, 114.05, 40.0),
        membership_audit_path=audit_path,
        prior_sampling_source=prior_sampling_path,
    )

    index_path = output_root / "china_full_1280m_macrocell_index.gpkg"
    assert index_path.exists()
    macrocells = gpd.read_file(index_path, layer="macrocells")
    utm_zones = gpd.read_file(index_path, layer="utm_zones")
    assert len(macrocells) == 2
    assert set(utm_zones["GRID_ID"]) == {"utm50n"}
    assert int(macrocells["ALL_COUNT"].sum()) == 12
    assert int(utm_zones["ALL_COUNT"].sum()) == 12
    assert int(macrocells["SAMPLED"].sum()) == 2
    assert int(macrocells["UNSAMPLED"].sum()) == 10

    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    assert manifest["counts"] == {"all": 12, "sampled": 2, "unsampled": 10}
    assert manifest["membership_audit"]["passed"] is True
    for preview in manifest["previews"].values():
        preview_path = output_root / preview
        assert preview_path.exists()
        assert preview_path.stat().st_size > 0
    assert (output_root / "先读我.docx").exists()


@pytest.mark.parametrize(
    ("audit", "message"),
    (
        (_valid_package_audit() | {"passed": False}, "passed must be exactly True"),
        (_valid_package_audit() | {"all_count": "12"}, "all_count must be a non-negative integer"),
        (
            {key: value for key, value in _valid_package_audit().items() if key != "matched"},
            "missing required fields: matched",
        ),
    ),
)
def test_package_artifacts_reject_invalid_audit_before_publishing(
    tmp_path: Path, audit: dict[str, object], message: str
) -> None:
    builder = _load_module(BUILDER_MODULE_PATH, "build_china_full_grid_invalid_audit")
    records = synthetic_parent_records(count=12)
    output_root = tmp_path / "package"
    MODULE.write_zone_records(
        records,
        {str(records[1]["parent_key"]), str(records[7]["parent_key"])},
        output_root,
        batch_size=5,
    )
    audit_path = tmp_path / "membership_audit.json"
    audit_path.write_text(json.dumps(audit), encoding="utf-8")

    with pytest.raises(ValueError, match=message):
        builder.build_package_artifacts(
            output_root=output_root,
            boundary_wgs84=box(113.95, 39.7, 114.05, 40.0),
            membership_audit_path=audit_path,
            prior_sampling_source=tmp_path / "unused.md",
        )

    _assert_package_artifacts_absent(output_root)


def test_writer_partitions_all_sampled_and_unsampled(tmp_path: Path) -> None:
    records = synthetic_parent_records(count=12)
    sampled = {str(records[1]["parent_key"]), str(records[7]["parent_key"])}

    summary = MODULE.write_zone_records(iter(records), sampled, tmp_path, batch_size=5)

    assert summary.all_count == 12
    assert summary.sampled_count == 2
    assert summary.unsampled_count == 10
    assert summary.sampled_count + summary.unsampled_count == summary.all_count

    def read_partition(name: str) -> gpd.GeoDataFrame:
        parts = sorted((tmp_path / name / "utm50n").glob("*.parquet"))
        frames = [gpd.read_parquet(part) for part in parts]
        return gpd.GeoDataFrame(pd.concat(frames), crs="EPSG:4326")

    all_records = read_partition("all")
    sampled_records = read_partition("sampled")
    unsampled_records = read_partition("unsampled")

    assert all_records.crs.to_epsg() == 4326
    assert set(sampled_records["parent_key"]).isdisjoint(unsampled_records["parent_key"])
    assert set(sampled_records["parent_key"]) == sampled
    assert len(all_records) == len(sampled_records) + len(unsampled_records)
    assert {"patch_id", "parent_key", "schema_version", "atlas_version", "boundary_version"} <= set(
        all_records.columns
    )
    assert all_records.geometry.geom_type.eq("Polygon").all()

    shapefile_parts = sorted((tmp_path / "all" / "utm50n").glob("*.shp"))
    shapefile_records = gpd.GeoDataFrame(
        pd.concat([gpd.read_file(part) for part in shapefile_parts]), crs="EPSG:4326"
    )
    assert shapefile_records.crs.to_epsg() == 4326
    assert len(shapefile_records) == 12
    assert {"PATCH_ID", "UTM_EPSG", "GRID_COL", "GRID_ROW", "MACRO_ID", "SAMPLED"} <= set(
        shapefile_records.columns
    )


def test_writer_rejects_shapefile_larger_than_cap(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setattr(MODULE, "SHAPEFILE_MAX_BYTES", 1)

    with pytest.raises(ValueError, match="one Shapefile feature exceeds"):
        MODULE.write_zone_records(synthetic_parent_records(count=1), set(), tmp_path, batch_size=1)

    assert not list(tmp_path.rglob("*.shp"))
    assert not list(tmp_path.rglob("*.parquet"))


def test_writer_splits_shapefiles_at_component_cap(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    records = synthetic_parent_records(count=12)
    probe = tmp_path / "probe.shp"
    MODULE._shapefile_records(records[:1], set()).to_file(probe, index=False)
    one_record_component = max(
        probe.with_suffix(suffix).stat().st_size for suffix in (".shp", ".shx", ".dbf")
    )
    component_cap = math.ceil(one_record_component / 0.9)
    monkeypatch.setattr(MODULE, "SHAPEFILE_MAX_BYTES", component_cap)

    MODULE.write_zone_records(records, set(), tmp_path / "output", batch_size=12)

    parts = sorted((tmp_path / "output" / "all" / "utm50n").glob("*.shp"))
    assert len(parts) > 1
    assert all("rowblock-" in part.name and "part-" in part.name for part in parts)
    for part in parts:
        assert all(
            part.with_suffix(suffix).stat().st_size < component_cap
            for suffix in (".shp", ".shx", ".dbf", ".prj", ".cpg")
            if part.with_suffix(suffix).exists()
        )


def test_geoparquet_has_canonical_fields_and_zstd_metadata(tmp_path: Path) -> None:
    record = synthetic_parent_records(count=1)[0]
    MODULE.write_zone_records([record], set(), tmp_path, batch_size=1)
    parquet_path = next((tmp_path / "all" / "utm50n").glob("*.parquet"))
    result = gpd.read_parquet(parquet_path)
    row = result.iloc[0]

    assert {"wgs84_bounds", "identity_hash", "footprint_hash"} <= set(result.columns)
    assert tuple(row["wgs84_bounds"]) == pytest.approx(row.geometry.bounds)
    expected_identity = (
        f"{row['atlas_version']}:{row['grid_epsg']}:{row['grid_col']}:{row['grid_row']}"
    )
    assert row["identity_hash"] == hashlib.sha256(expected_identity.encode("utf-8")).hexdigest()
    expected_footprint = normalize(set_precision(row.geometry, 1e-9)).wkb
    assert row["footprint_hash"] == hashlib.sha256(expected_footprint).hexdigest()

    metadata = pq.ParquetFile(parquet_path).metadata
    assert {
        metadata.row_group(0).column(column).compression
        for column in range(metadata.row_group(0).num_columns)
    } == {"ZSTD"}


def test_membership_audit_requires_every_sample_exactly_once() -> None:
    atlas = ["32650:1:1", "32650:1:2", "32650:1:3"]
    sampled = ["32650:1:1", "32650:1:3"]

    audit = MODULE.audit_sample_membership(atlas, sampled)

    assert audit["all_count"] == 3
    assert audit["sampled_count"] == 2
    assert audit["unsampled_count"] == 1
    assert audit["matched"] == 2
    assert audit["missing"] == []
    assert audit["duplicate_atlas_keys"] == []


def test_membership_audit_reports_missing_and_duplicated_sampled_cells() -> None:
    atlas = ["32650:1:1", "32650:1:1", "32650:1:2"]
    sampled = ["32650:1:1", "32650:1:3", "32650:1:3"]

    audit = MODULE.audit_sample_membership(atlas, sampled)

    assert audit["matched"] == 0
    assert audit["missing"] == ["32650:1:3"]
    assert audit["duplicate_atlas_keys"] == ["32650:1:1"]
    assert audit["duplicate_sampled_keys"] == ["32650:1:3"]


def test_partitioned_audit_detects_altered_footprints_and_cross_zone_overlap(
    tmp_path: Path,
) -> None:
    records = synthetic_parent_records(count=2)
    sampled_registry = []
    for record in records:
        geometry = MODULE._wgs84_geometry(record)
        sampled_registry.append(
            {
                "grid_epsg": record["grid_epsg"],
                "grid_col": record["grid_col"],
                "grid_row": record["grid_row"],
                "canonical_wgs84_footprint_hash": hashlib.sha256(
                    normalize(set_precision(geometry, 1e-9)).wkb
                ).hexdigest(),
            }
        )
    MODULE.write_zone_records(records, {str(records[0]["parent_key"])}, tmp_path, batch_size=1)

    parquet_path = next((tmp_path / "all" / "utm50n").glob("*.parquet"))
    frame = gpd.read_parquet(parquet_path)
    frame.loc[0, "footprint_hash"] = "0" * 64
    frame.loc[0, "geometry"] = frame.loc[0, "geometry"].buffer(0.0001)
    frame.to_parquet(parquet_path, index=False, compression="zstd")

    audit = MODULE.audit_grid_package(tmp_path, sampled_registry, batch_size=1)

    assert audit["all_count"] == 2
    assert audit["sampled_count"] == 1
    assert audit["unsampled_count"] == 1
    assert audit["hash_mismatches"]["footprint_hash"] == 1
    assert audit["hash_mismatches"]["sampled_registry_footprint_hash"] == 1
    assert audit["max_footprint_coordinate_difference"] > 0
    assert audit["invalid_geometry_count"] == 1


def test_partitioned_audit_reports_same_and_cross_zone_positive_area_overlaps(
    tmp_path: Path,
) -> None:
    first = synthetic_parent_records(count=1)[0]
    to_utm49 = Transformer.from_crs(4326, 32649, always_xy=True)
    to_wgs84 = Transformer.from_crs(32649, 4326, always_xy=True)
    easting, northing = to_utm49.transform(111.0, 39.84)
    grid_col = math.floor(easting / MODULE.PARENT_SIDE_METERS)
    grid_row = math.floor(northing / MODULE.PARENT_SIDE_METERS)
    longitude, latitude = to_wgs84.transform(
        (grid_col + 0.5) * MODULE.PARENT_SIDE_METERS,
        (grid_row + 0.5) * MODULE.PARENT_SIDE_METERS,
    )
    second = MODULE.build_patch_record(
        {
            "grid_epsg": 32649,
            "grid_id": "utm49n",
            "macro_col": grid_col // 10,
            "macro_row": grid_row // 10,
        },
        grid_col,
        grid_row,
        longitude,
        latitude,
        MODULE.GridSpec(boundary_version="test"),
    )
    records = [first, second]
    for record in records:
        MODULE.write_zone_records([record], set(), tmp_path, batch_size=1)

    first_path = next((tmp_path / "all" / "utm50n").glob("*.parquet"))
    second_path = next((tmp_path / "all" / "utm49n").glob("*.parquet"))
    first_frame = gpd.read_parquet(first_path)
    second_frame = gpd.read_parquet(second_path)
    second_frame.loc[0, "geometry"] = first_frame.loc[0, "geometry"]
    second_frame.to_parquet(second_path, index=False, compression="zstd")

    audit = MODULE.audit_grid_package(tmp_path, [], batch_size=1)

    assert audit["same_zone_positive_overlap_count"] == 0
    assert audit["cross_zone_overlap_violation_count"] == 1
    assert audit["max_cross_zone_overlap_fraction"] == pytest.approx(1.0)


def test_partitioned_audit_reports_same_zone_positive_area_overlap(tmp_path: Path) -> None:
    records = synthetic_parent_records(count=2)
    MODULE.write_zone_records(records, set(), tmp_path, batch_size=1)
    paths = sorted((tmp_path / "all" / "utm50n").glob("*.parquet"))
    frame = gpd.read_parquet(paths[0])
    frame.loc[1, "geometry"] = frame.loc[0, "geometry"]
    frame.to_parquet(paths[0], index=False, compression="zstd")

    audit = MODULE.audit_grid_package(tmp_path, [], batch_size=1)

    assert audit["same_zone_positive_overlap_count"] == 1


def test_partitioned_audit_rejects_shifted_footprint_with_recomputed_hash(tmp_path: Path) -> None:
    record = synthetic_parent_records(count=1)[0]
    MODULE.write_zone_records([record], set(), tmp_path, batch_size=1)
    parquet_path = next((tmp_path / "all" / "utm50n").glob("*.parquet"))
    frame = gpd.read_parquet(parquet_path)
    shifted_geometry = translate(frame.loc[0, "geometry"], xoff=0.0001)
    frame.loc[0, "geometry"] = shifted_geometry
    frame.loc[0, "footprint_hash"] = hashlib.sha256(
        normalize(set_precision(shifted_geometry, 1e-9)).wkb
    ).hexdigest()
    frame.to_parquet(parquet_path, index=False, compression="zstd")

    audit = MODULE.audit_grid_package(tmp_path, [], batch_size=1)

    assert audit["hash_mismatches"]["footprint_hash"] == 1
    assert audit["footprint_coordinate_mismatch_count"] == 1
    assert audit["passed"] is False


def test_partitioned_audit_derives_all_footprint_from_integer_grid_coordinates(
    tmp_path: Path,
) -> None:
    record = synthetic_parent_records(count=1)[0]
    MODULE.write_zone_records([record], set(), tmp_path, batch_size=1)
    parquet_path = next((tmp_path / "all" / "utm50n").glob("*.parquet"))
    frame = gpd.read_parquet(parquet_path)
    shifted_bounds = [value + MODULE.PARENT_SIDE_METERS for value in frame.loc[0, "utm_bounds"]]
    shifted_geometry = MODULE._wgs84_geometry(
        {"grid_epsg": frame.loc[0, "grid_epsg"], "utm_bounds": shifted_bounds}
    )
    frame.at[0, "utm_bounds"] = shifted_bounds
    frame.at[0, "wgs84_bounds"] = list(shifted_geometry.bounds)
    frame.loc[0, "geometry"] = shifted_geometry
    frame.loc[0, "footprint_hash"] = hashlib.sha256(
        normalize(set_precision(shifted_geometry, 1e-9)).wkb
    ).hexdigest()
    frame.to_parquet(parquet_path, index=False, compression="zstd")

    audit = MODULE.audit_grid_package(tmp_path, [], batch_size=1)

    assert audit["stored_utm_bounds_mismatch_count"] == 1
    assert audit["hash_mismatches"]["footprint_hash"] == 1
    assert audit["passed"] is False


def test_partitioned_audit_rejects_swapped_sampled_and_unsampled_children(tmp_path: Path) -> None:
    records = synthetic_parent_records(count=2)
    sampled_key = str(records[0]["parent_key"])
    MODULE.write_zone_records(records, {sampled_key}, tmp_path, batch_size=2)
    sampled_path = next((tmp_path / "sampled" / "utm50n").glob("*.parquet"))
    unsampled_path = next((tmp_path / "unsampled" / "utm50n").glob("*.parquet"))
    sampled_payload = sampled_path.read_bytes()
    unsampled_payload = unsampled_path.read_bytes()
    sampled_path.write_bytes(unsampled_payload)
    unsampled_path.write_bytes(sampled_payload)

    audit = MODULE.audit_grid_package(
        tmp_path,
        [
            {
                "grid_epsg": records[0]["grid_epsg"],
                "grid_col": records[0]["grid_col"],
                "grid_row": records[0]["grid_row"],
            }
        ],
        batch_size=1,
    )

    assert audit["exact_partition_membership_mismatch_count"] == 2
    assert audit["passed"] is False


def test_partitioned_audit_rejects_rehashed_child_metadata_and_geometry(tmp_path: Path) -> None:
    record = synthetic_parent_records(count=1)[0]
    MODULE.write_zone_records([record], {str(record["parent_key"])}, tmp_path, batch_size=1)
    parquet_path = next((tmp_path / "sampled" / "utm50n").glob("*.parquet"))
    frame = gpd.read_parquet(parquet_path)
    shifted_geometry = translate(frame.loc[0, "geometry"], xoff=0.0001)
    frame.loc[0, "geometry"] = shifted_geometry
    frame.at[0, "wgs84_bounds"] = list(shifted_geometry.bounds)
    frame.loc[0, "footprint_hash"] = hashlib.sha256(
        normalize(set_precision(shifted_geometry, 1e-9)).wkb
    ).hexdigest()
    frame.to_parquet(parquet_path, index=False, compression="zstd")

    audit = MODULE.audit_grid_package(
        tmp_path,
        [
            {
                "grid_epsg": record["grid_epsg"],
                "grid_col": record["grid_col"],
                "grid_row": record["grid_row"],
            }
        ],
        batch_size=1,
    )

    assert audit["child_partition_metadata_mismatch_count"] == 1
    assert audit["child_partition_geometry_mismatch_count"] == 1
    assert audit["passed"] is False


def test_cli_writes_json_membership_audit_from_jsonl_registry(tmp_path: Path) -> None:
    records = synthetic_parent_records(count=2)
    output_root = tmp_path / "output"
    registry_path = tmp_path / "sampled-registry.jsonl"
    audit_path = tmp_path / "china_full_grid_membership_audit.json"
    MODULE.write_zone_records(records, {str(records[0]["parent_key"])}, output_root, batch_size=1)
    registry_path.write_text(
        json.dumps(
            {
                "grid_epsg": records[0]["grid_epsg"],
                "grid_col": records[0]["grid_col"],
                "grid_row": records[0]["grid_row"],
            }
        )
        + "\n",
        encoding="utf-8",
    )

    result = subprocess.run(
        [
            sys.executable,
            "scripts/data/build_china_full_grid.py",
            "--audit-only",
            "--output-root",
            str(output_root),
            "--sampled-registry",
            str(registry_path),
            "--audit-output",
            str(audit_path),
            "--batch-size",
            "1",
        ],
        cwd=MODULE_PATH.parents[2],
        capture_output=True,
        check=False,
        text=True,
    )

    assert result.returncode == 0, result.stderr
    assert audit_path.exists()
    audit = json.loads(audit_path.read_text(encoding="utf-8"))
    assert audit["passed"] is True
    assert audit["matched"] == 1


def test_writer_keeps_output_empty_when_a_later_partition_write_fails(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    records = synthetic_parent_records(count=2)
    sampled = {str(records[0]["parent_key"])}
    original_to_parquet = gpd.GeoDataFrame.to_parquet
    calls = 0

    def fail_second_parquet_write(self, path, *args, **kwargs):
        nonlocal calls
        calls += 1
        if calls == 2:
            raise OSError("synthetic second partition failure")
        return original_to_parquet(self, path, *args, **kwargs)

    monkeypatch.setattr(gpd.GeoDataFrame, "to_parquet", fail_second_parquet_write)

    with pytest.raises(OSError, match="synthetic second partition failure"):
        MODULE.write_zone_records(records, sampled, tmp_path, batch_size=2)

    assert not list(tmp_path.rglob("*.parquet"))
    assert not list(tmp_path.rglob("*.shp"))


def test_cli_writes_requested_zone_from_small_local_inputs(tmp_path: Path) -> None:
    record = next(MODULE.enumerate_macro_patch_records(_macro(), _boundary(), MODULE.GridSpec()))
    boundary_path = tmp_path / "boundary.geojson"
    macro_path = tmp_path / "macros.json"
    sampled_path = tmp_path / "sampled.json"
    output_root = tmp_path / "output"
    boundary_frame = gpd.GeoDataFrame(geometry=[_boundary()], crs="EPSG:4326")
    boundary_frame.to_file(boundary_path, driver="GeoJSON")
    macro_path.write_text(
        '[{"grid_epsg": 32650, "grid_id": "utm50n", "macro_col": 34, "macro_row": 344}]',
        encoding="utf-8",
    )
    sampled_key = record["parent_key"]
    sampled_path.write_text(f'["{sampled_key}"]', encoding="utf-8")

    result = subprocess.run(
        [
            sys.executable,
            "scripts/data/build_china_full_grid.py",
            "--boundary",
            str(boundary_path),
            "--macrocells",
            str(macro_path),
            "--sampled-parent-keys",
            str(sampled_path),
            "--output-root",
            str(output_root),
            "--grid-id",
            "utm50n",
            "--batch-size",
            "100",
        ],
        cwd=MODULE_PATH.parents[2],
        capture_output=True,
        check=False,
        text=True,
    )

    assert result.returncode == 0, result.stderr
    assert '"sampled_count": 1' in result.stdout
    assert list((output_root / "all" / "utm50n").glob("*.shp"))


def test_grid_spec_rejects_noncanonical_parent_cell_size() -> None:
    with pytest.raises(ValueError, match="side_m must be exactly 1280"):
        MODULE.GridSpec(side_m=640)


def test_enumerator_emits_exact_1280m_parent_cells() -> None:
    spec = MODULE.GridSpec(side_m=1280, macro_side_patches=10, boundary_version="test")
    records = list(MODULE.enumerate_macro_patch_records(_macro(), _boundary(), spec))

    assert records
    for record in records:
        minx, miny, maxx, maxy = record["utm_bounds"]
        assert maxx - minx == 1280
        assert maxy - miny == 1280
        assert minx == record["grid_col"] * 1280
        assert miny == record["grid_row"] * 1280
        assert record["parent_key"] == (
            f'{record["grid_epsg"]}:{record["grid_col"]}:{record["grid_row"]}'
        )


def test_enumerator_records_have_deterministic_ids() -> None:
    spec = MODULE.GridSpec(boundary_version="test")

    first = list(MODULE.enumerate_macro_patch_records(_macro(), _boundary(), spec))
    second = list(MODULE.enumerate_macro_patch_records(_macro(), _boundary(), spec))

    assert [record["patch_id"] for record in first] == [record["patch_id"] for record in second]
    assert all(record["patch_id"] == f"parent_{record['parent_key']}" for record in first)


def test_enumerator_gives_each_record_a_unique_macro_local_key() -> None:
    spec = MODULE.GridSpec(boundary_version="test")
    records = list(MODULE.enumerate_macro_patch_records(_macro(), _boundary(), spec))

    keys = {
        (record["macro_id"], record["macro_local_col"], record["macro_local_row"])
        for record in records
    }
    assert len(keys) == len(records)
    assert all(0 <= record["macro_local_col"] < 10 for record in records)
    assert all(0 <= record["macro_local_row"] < 10 for record in records)


def test_enumerator_filters_cells_owned_by_another_utm_zone() -> None:
    spec = MODULE.GridSpec(boundary_version="test")
    to_utm49 = Transformer.from_crs(4326, 32649, always_xy=True)
    easting, northing = to_utm49.transform(117.0, 39.84)
    macro = {
        "grid_epsg": 32649,
        "grid_id": "utm49n",
        "macro_col": math.floor(easting / 12800),
        "macro_row": math.floor(northing / 12800),
    }
    boundary = box(116.9, 39.7, 117.1, 40.0)

    records = list(MODULE.enumerate_macro_patch_records(macro, boundary, spec))

    assert records == []


def test_enumerator_rejects_epsg_outside_china_owner_range() -> None:
    macro = {**_macro(), "grid_epsg": 32642, "grid_id": "utm42n"}

    with pytest.raises(ValueError, match="32643 through 32653"):
        list(MODULE.enumerate_macro_patch_records(macro, _boundary(), MODULE.GridSpec()))


def test_enumerator_preserves_wgs84_longitude_latitude_axis_order() -> None:
    spec = MODULE.GridSpec(boundary_version="test")
    records = list(MODULE.enumerate_macro_patch_records(_macro(), _boundary(), spec))

    assert records
    for record in records:
        assert 116.28 <= record["longitude"] <= 116.31
        assert 39.83 <= record["latitude"] <= 39.86


def test_validate_patch_record_rejects_misaligned_bounds() -> None:
    spec = MODULE.GridSpec(boundary_version="test")
    record = next(MODULE.enumerate_macro_patch_records(_macro(), _boundary(), spec))
    record["utm_bounds"][0] += 1

    with pytest.raises(ValueError, match="utm_bounds"):
        MODULE.validate_patch_record(record, spec)


def test_validate_patch_record_rejects_epsg_outside_china_owner_range() -> None:
    spec = MODULE.GridSpec(boundary_version="test")
    to_utm42 = Transformer.from_crs(4326, 32642, always_xy=True)
    easting, northing = to_utm42.transform(69.0, 30.0)
    grid_col = math.floor(easting / 1280)
    grid_row = math.floor(northing / 1280)
    macro = {
        "grid_epsg": 32642,
        "grid_id": "utm42n",
        "macro_col": grid_col // 10,
        "macro_row": grid_row // 10,
    }
    record = MODULE.build_patch_record(macro, grid_col, grid_row, 69.0, 30.0, spec)

    with pytest.raises(ValueError, match="32643 through 32653"):
        MODULE.validate_patch_record(record, spec)
