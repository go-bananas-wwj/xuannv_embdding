from __future__ import annotations

import importlib.util
import sys
from pathlib import Path

from PIL import Image
from pptx import Presentation

SCRIPT = Path(__file__).parents[1] / "scripts/report/build_pujiang_slide02.py"
SPEC = importlib.util.spec_from_file_location("build_pujiang_slide02", SCRIPT)
MODULE = importlib.util.module_from_spec(SPEC)
assert SPEC.loader is not None
sys.modules[SPEC.name] = MODULE
SPEC.loader.exec_module(MODULE)


def test_build_monthly_input_montage_contains_three_sources_and_six_months(
    tmp_path: Path,
) -> None:
    source = tmp_path / "qa.png"
    Image.new("RGB", (3230, 2040), "white").save(source)
    output = tmp_path / "montage.png"

    metadata = MODULE.build_monthly_input_montage(source, output)

    assert output.exists()
    assert metadata["months"] == [
        "2025.12",
        "2026.01",
        "2026.02",
        "2026.03",
        "2026.04",
        "2026.05",
    ]
    assert metadata["sources"] == ["Sentinel-2", "高分光学", "高分SAR"]
    with Image.open(output) as image:
        assert image.size == (1800, 630)


def test_slide02_contains_evidence_bounded_training_facts(tmp_path: Path) -> None:
    montage = tmp_path / "montage.png"
    Image.new("RGB", (1800, 630), "white").save(montage)
    output = tmp_path / "slide02.pptx"

    MODULE.build_ppt(MODULE.DEFAULT_SOURCE_PPTX, output, montage)

    presentation = Presentation(output)
    assert len(presentation.slides) == 1
    text = "\n".join(
        shape.text for shape in presentation.slides[0].shapes if hasattr(shape, "text")
    )
    for expected in [
        "左侧流程与样例：海淀 P10C",
        "区域网格划分",
        "云雾与无效像素筛选",
        "像素级 mask",
        "320 patch",
        "P10C epoch 800",
        "2×Ascend NPU",
        "约 17 小时 50 分",
        "从 P9B epoch 800 继续训练",
        "424 patch",
        "V5（API v2）",
        "训练数据",
        "API 未披露",
        "训练硬件",
        "训练时长",
        "未混用旧 Stage 2 训练记录",
    ]:
        assert expected in text


def test_preview_has_widescreen_dimensions(tmp_path: Path) -> None:
    montage = tmp_path / "montage.png"
    Image.new("RGB", (1800, 630), "white").save(montage)
    output = tmp_path / "preview.png"

    MODULE.build_preview(output, montage)

    with Image.open(output) as image:
        assert image.size == (1600, 900)
