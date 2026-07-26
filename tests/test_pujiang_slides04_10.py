from __future__ import annotations

import importlib.util
import sys
from pathlib import Path

from PIL import Image
from pptx import Presentation

SCRIPT = Path(__file__).parents[1] / "scripts/report/build_pujiang_slides04_10.py"
SPEC = importlib.util.spec_from_file_location("build_pujiang_slides04_10", SCRIPT)
MODULE = importlib.util.module_from_spec(SPEC)
assert SPEC.loader is not None
sys.modules[SPEC.name] = MODULE
SPEC.loader.exec_module(MODULE)


def test_slide04_builds_platform_overview(tmp_path: Path) -> None:
    output = tmp_path / "slide04.pptx"
    preview = tmp_path / "slide04.png"

    metadata = MODULE.build_page(4, output, preview)

    assert len(Presentation(output).slides) == 1
    assert metadata["title"] == "玄女嵌入平台：浏览、分析、训练与生成"
    assert metadata["functions"] == ["嵌入浏览", "专题分析", "自定义训练", "时相生成"]
    with Image.open(preview) as image:
        assert image.size == (1600, 900)


def test_slides05_to10_build_with_required_evidence(tmp_path: Path) -> None:
    expectations = {
        5: ("少量标注，快速形成区域级专题图", "5-shot"),
        6: ("嵌入底座进入遥感智能体工作流", "embedding_tool"),
        7: ("云遮挡或观测缺失时，生成指定时刻的遥感参考影像", "参考影像"),
        8: ("海淀区同协议评测：64 维月度嵌入达到 AEF 同量级", 0.484),
        9: ("标注越少，月度嵌入的复用价值越明显", 41.5),
        10: ("哈尔滨新区：从月度嵌入到城市治理专题", "V5/v2 API"),
    }
    for page, (title, evidence) in expectations.items():
        output = tmp_path / f"slide{page:02d}.pptx"
        preview = tmp_path / f"slide{page:02d}.png"

        metadata = MODULE.build_page(page, output, preview)

        assert len(Presentation(output).slides) == 1
        assert metadata["title"] == title
        assert str(evidence) in str(metadata)
        with Image.open(preview) as image:
            assert image.size == (1600, 900)
