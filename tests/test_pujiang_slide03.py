from __future__ import annotations

import importlib.util
import sys
from pathlib import Path

from PIL import Image
from pptx import Presentation

SCRIPT = Path(__file__).parents[1] / "scripts/report/build_pujiang_slide03.py"
SPEC = importlib.util.spec_from_file_location("build_pujiang_slide03", SCRIPT)
MODULE = importlib.util.module_from_spec(SPEC)
assert SPEC.loader is not None
sys.modules[SPEC.name] = MODULE
SPEC.loader.exec_module(MODULE)


def test_slide03_states_p10c_adaptations_without_overclaiming(tmp_path: Path) -> None:
    output = tmp_path / "slide03.pptx"

    MODULE.build_ppt(MODULE.DEFAULT_SOURCE_PPTX, output)

    presentation = Presentation(output)
    assert len(presentation.slides) == 1
    text = "\n".join(
        shape.text for shape in presentation.slides[0].shapes if hasattr(shape, "text")
    )
    for expected in [
        "参考 AEF 的 embedding field 范式",
        "不是完全复现 AEF",
        "128×128×64",
        "连续 6 期",
        "跨月份注意力",
        "先故意拿走一部分输入",
        "整源模态遮挡",
        "S2 18%",
        "S1 / Landsat 35%",
        "每个训练样本有 65% 概率启动月份遮挡",
        "随机拿走 1–4 个月",
        "每个训练样本有 65% 概率启动空间遮挡",
        "遮住约 32% 的图像块",
        "答案仍是完整影像",
        "从剩余月份、其他传感器和周边空间恢复",
        "availability mask",
        "只记录原始数据是否真实存在",
        "OSM 只做语义提示",
        "13 项可重叠的 OSM 语义提示",
        "建筑、主路、支路、铁路、水体、绿地、农业、居住、商业、工业、施工、步道、操场",
        "对每个像素分别判断是否属于各项可重叠语义",
        "0.14 控制整组 OSM 提示对训练的影响力度",
        "不代表总损失的 14%",
        "临时 1×1 线性探针",
        "训练结束后丢弃",
        "前 80 轮逐步加到 0.14",
        "人工变化监督 0",
        "时间对比损失 0",
    ]:
        assert expected in text
    font_sizes = [
        run.font.size.pt
        for shape in presentation.slides[0].shapes
        if hasattr(shape, "text_frame")
        for paragraph in shape.text_frame.paragraphs
        for run in paragraph.runs
        if run.text.strip() and run.font.size is not None
    ]
    assert min(font_sizes) >= 12


def test_preview_has_widescreen_dimensions(tmp_path: Path) -> None:
    output = tmp_path / "preview.png"

    MODULE.build_preview(output)

    with Image.open(output) as image:
        assert image.size == (1600, 900)
