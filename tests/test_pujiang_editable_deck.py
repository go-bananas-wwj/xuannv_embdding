from __future__ import annotations

import importlib.util
import sys
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Inches

SCRIPT = Path(__file__).parents[1] / "scripts/report/pujiang_editable_common.py"
SPEC = importlib.util.spec_from_file_location("pujiang_editable_common", SCRIPT)
MODULE = importlib.util.module_from_spec(SPEC)
assert SPEC.loader is not None
sys.modules[SPEC.name] = MODULE
SPEC.loader.exec_module(MODULE)


def test_helpers_create_separate_text_and_picture_shapes(tmp_path: Path) -> None:
    image_path = tmp_path / "observation.png"
    Image.new("RGB", (80, 40), (33, 128, 94)).save(image_path)
    presentation = Presentation()
    presentation.slide_width = Inches(13.333333)
    presentation.slide_height = Inches(7.5)

    slide = MODULE.new_blank_slide(presentation)
    text_shape = MODULE.add_text(slide, "可编辑标题", 0.6, 0.7, 3.2, 0.5, 22)
    picture_shape = MODULE.add_picture(slide, image_path, 4.0, 1.0, 2.4, 1.2)

    assert text_shape is not picture_shape
    assert text_shape.text == "可编辑标题"
    assert picture_shape.shape_type == MSO_SHAPE_TYPE.PICTURE
    assert len(slide.shapes) == 2


def test_full_slide_picture_detector_uses_presentation_geometry(tmp_path: Path) -> None:
    image_path = tmp_path / "background.png"
    Image.new("RGB", (160, 90), (236, 246, 252)).save(image_path)
    presentation = Presentation()
    presentation.slide_width = Inches(13.333333)
    presentation.slide_height = Inches(7.5)
    slide = MODULE.new_blank_slide(presentation)

    full_slide_picture = slide.shapes.add_picture(
        str(image_path),
        0,
        0,
        width=presentation.slide_width,
        height=presentation.slide_height,
    )
    inset_picture = MODULE.add_picture(slide, image_path, 0.6, 0.7, 3.2, 1.8)

    assert MODULE.is_full_slide_picture(full_slide_picture, presentation)
    assert not MODULE.is_full_slide_picture(inset_picture, presentation)


def test_clone_first_slide_keeps_picture_relationships(tmp_path: Path) -> None:
    image_path = tmp_path / "source.png"
    source_path = tmp_path / "source.pptx"
    destination_path = tmp_path / "destination.pptx"
    Image.new("RGB", (80, 40), (0, 97, 170)).save(image_path)

    source = Presentation()
    source.slide_width = Inches(13.333333)
    source.slide_height = Inches(7.5)
    source_slide = MODULE.new_blank_slide(source)
    MODULE.add_text(source_slide, "原生文字", 0.6, 0.7, 3.2, 0.5, 22)
    MODULE.add_picture(source_slide, image_path, 4.0, 1.0, 2.4, 1.2)
    source.save(source_path)

    destination = Presentation()
    MODULE.clone_first_slide(source_path, destination)
    destination.save(destination_path)

    cloned = Presentation(destination_path).slides[0]
    assert any(shape.text == "原生文字" for shape in cloned.shapes if hasattr(shape, "text"))
    assert any(shape.shape_type == MSO_SHAPE_TYPE.PICTURE for shape in cloned.shapes)
