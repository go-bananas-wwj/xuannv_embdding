from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


ROOT = Path(__file__).resolve().parents[2]
SOURCE = ROOT / "docs/projects/haidian_gep_xuannv_implementation_and_acceptance_annex_20260729_zh.md"
OUTPUT = ROOT / "docs/projects/haidian_gep_xuannv_implementation_and_acceptance_annex_20260729_zh.docx"


def set_run_font(run, size: Pt, bold: bool = False) -> None:
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    run.font.size = size
    run.font.bold = bold
    run.font.color.rgb = None


def add_text(paragraph, text: str, size: Pt, bold: bool = False) -> None:
    parts = re.split(r"(\*\*[^*]+\*\*)", text)
    for part in parts:
        is_bold = part.startswith("**") and part.endswith("**")
        run = paragraph.add_run(part[2:-2] if is_bold else part)
        set_run_font(run, size, bold or is_bold)


def set_cell_text(cell, text: str, *, bold: bool = False) -> None:
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.space_before = Pt(0)
    add_text(p, text, Pt(10.5), bold=bold)
    tc_pr = cell._tc.get_or_add_tcPr()
    v_align = OxmlElement("w:vAlign")
    v_align.set(qn("w:val"), "center")
    tc_pr.append(v_align)


def set_table_borders(table) -> None:
    tbl_pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = OxmlElement(f"w:{edge}")
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), "6")
        tag.set(qn("w:color"), "000000")
        borders.append(tag)
    tbl_pr.append(borders)


def add_info_table(document: Document, project_title: str) -> None:
    table = document.add_table(rows=5, cols=2)
    table.autofit = False
    table.columns[0].width = Cm(4.0)
    table.columns[1].width = Cm(12.5)
    entries = [
        ("项目名称", project_title),
        ("申报部门", "待填"),
        ("项目申报人", "待填"),
        ("研究周期", "建议2年（以正式立项文件为准）"),
        ("成果转化方式", "业务化、工程化"),
    ]
    for row, (label, value) in zip(table.rows, entries, strict=True):
        set_cell_text(row.cells[0], label, bold=True)
        set_cell_text(row.cells[1], value)
    set_table_borders(table)


def build_document() -> None:
    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    document = Document()
    section = document.sections[0]
    section.top_margin = Cm(2.54)
    section.bottom_margin = Cm(2.54)
    section.left_margin = Cm(2.7)
    section.right_margin = Cm(2.7)

    style = document.styles["Normal"]
    style.font.name = "Times New Roman"
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    style.font.size = Pt(12)
    style.paragraph_format.line_spacing = 1.5
    style.paragraph_format.first_line_indent = Cm(0.74)

    title = lines[0].removeprefix("# ").strip()
    title_p = document.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_p.paragraph_format.space_after = Pt(28)
    add_text(title_p, title, Pt(16), bold=True)
    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(42)
    add_text(subtitle, "项目申报书（技术内容修订稿）", Pt(14), bold=True)
    document.add_paragraph("说明：申报部门、负责人、研究周期和经费预算需由申报单位补充。")
    document.add_page_break()
    add_info_table(document, title)
    document.add_page_break()

    in_info = False
    for raw in lines[1:]:
        line = raw.strip()
        if not line:
            continue
        if line == "## 项目申报信息":
            in_info = True
            continue
        if in_info:
            if line.startswith("## "):
                in_info = False
            else:
                continue
        if line.startswith("## "):
            p = document.add_paragraph()
            p.paragraph_format.space_before = Pt(14)
            p.paragraph_format.space_after = Pt(8)
            add_text(p, line[3:], Pt(14), bold=True)
        elif line.startswith("### "):
            p = document.add_paragraph()
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(4)
            add_text(p, line[4:], Pt(12), bold=True)
        else:
            p = document.add_paragraph()
            p.paragraph_format.space_after = Pt(5)
            p.paragraph_format.line_spacing = 1.5
            p.paragraph_format.first_line_indent = Cm(0.74)
            add_text(p, line, Pt(12))

    for section in document.sections:
        section.header_distance = Cm(1.5)
        footer = section.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_text(footer, "海淀区GEP核算与智能服务关键技术研究", Pt(9))

    document.save(OUTPUT)


if __name__ == "__main__":
    build_document()
