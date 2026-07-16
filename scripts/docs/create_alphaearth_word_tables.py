#!/usr/bin/env python3
"""Generate native Word/HTML tables for the China AlphaEarth planning review."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt


ROOT = Path("/root/workspace/xuannv/docs/reviews")
DOCX_OUT = ROOT / "china_alphaearth_word_copy_tables_20260716.docx"
HTML_OUT = ROOT / "china_alphaearth_word_copy_tables_20260716.html"


AUDIT_HEADERS = ["项目", "海淀区当前实测", "全国 30% / 6 个月估算", "统计口径与备注"]
AUDIT_ROWS = [
    [
        "原始影像与基础标签",
        "127.09 GiB（约 128 G）\nS2：约 104 G\n高分光学：约 11 G\nS1：约 9.2 G\nLandsat：约 3.5 G",
        "不宜直接按 127.09 GiB 线性外推。\n全国规模取决于每月保留帧数、高分数据覆盖、云筛选和原始归档策略。",
        "仅统计 raw/haidian。当前 raw 目录中未保留原始高分 SAR；不能将该目录视为完整全国原始数据密度。",
    ],
    [
        "处理后训练数据",
        "37.84 GiB（约 39 G）\n其中 patch 数据约 39 G，标签约 121 MiB",
        "当前 P10C 筛选训练包密度外推约 71 TB，约 3.11 亿文件。\n12 个月约 143 TB。",
        "全国训练必须转换为 512 MiB–2 GiB shard（WebDataset/tar、Zarr 等），不能直接随机读取数亿小文件。",
    ],
    [
        "P10C 生产版月度 embedding",
        "7.51 GiB（约 7.6 G）\n320 patch × 6 月\n当前每 patch×月约 4 MiB，FP32 保存",
        "约 44 TB（当前 FP32 导出）",
        "当前真实导出格式，不是理论 FP16 值。模型输出为 64 维、128×128 embedding map。",
    ],
    [
        "FP16 embedding（待验证优化）",
        "当前未采用",
        "约 22 TB",
        "仅在少样本制图、检索和变化任务验证精度无明显下降后，才可作为发布格式。",
    ],
    [
        "P10C checkpoint、日志与训练指标",
        "约 1.6 G",
        "模型权重不是全国容量主项；建议保留 best、last 和可复现训练配置。",
        "对应 outputs/v2_p10c...，不包括所有历史版本。",
    ],
    [
        "海淀生产工作集小计",
        "约 174 GiB",
        "建议至少 0.5 PB 可用存储，1 PB 原始容量",
        "小计 = 原始数据 + 处理数据 + 当前生产 embedding + P10C 训练产物；不含重复发布包。",
    ],
    [
        "历史/上传重复物理占用",
        "至少约 211 GiB 额外占用\n海淀相关全盘物理占用至少约 385 GiB",
        "不纳入全国生产主容量",
        "包括约 25 G 发布训练包、约 99 G ModelScope 暂存及旧 embedding、旧 checkpoint、历史可视化。部分目录为海淀/哈尔滨混合，故为保守上界。",
    ],
    [
        "5 m / 2 m 当前 FP32 embedding",
        "未生成全国产品",
        "约 177 TB / 1.11 PB",
        "分别约为 10 m 的 4 倍 / 25 倍；还未计全国高分输入数据、索引、金字塔和副本。",
    ],
]

AEF_HEADERS = ["对比项", "AlphaEarth Foundations（公开信息）", "玄女海淀 P10C 当前实测", "解释与启示"]
AEF_ROWS = [
    ["训练地点/覆盖", "约 841 万唯一时空地点，约地球陆地面积的 1.1%", "320 patch，约 524 km²", "AEF 训练地点规模远大于城市试点；中国版应先做 5–10 万 patch 的全国分层试点。"],
    ["训练观测量", "约 30.48 亿图像帧；9 类栅格数据源 + 1 类非结构化文本源", "6 个月 S2、S1、Landsat、高分光学/SAR、OSM 弱语义", "应按有效观测帧、时相、云质量和缺测率，而非仅面积预算。"],
    ["最终模型规模", "训练约 10 亿和约 4.8 亿参数版本；最终采用约 4.8 亿版本", "约 3,390 万参数量级", "不同模型规模和硬件代际，不能将 device-hours 直接换算成卡数或成本。"],
    ["训练算力", "512 TPU v4，56 小时，100,000 steps\n合计约 28,672 TPU v4 device-hours", "6 Ascend 910B，17 小时 49 分 32 秒，约 7,100 updates\n合计约 107 NPU-hours", "两者只能作为规模参照。中国版必须用自身完整配方实测 step/s 和扩展效率。"],
    ["Batch / 序列", "batch 256 个视频序列；每序列抽样 103 帧\n名义约 26,368 帧/step", "每卡 batch 3、梯度累积 2；6 卡有效全局 batch 约 36", "全国训练必须写固定 updates、全局 batch、有效帧/patch 和实测吞吐，不应只写 epoch。"],
    ["公开 embedding 产品", "年度、约 10 m、64 维；全球 2017–2025", "月度、10 m、64 维；2025-12 至 2026-05", "月度更新是玄女的差异化方向，但必须处理月度云雾、观测稀疏和缺失模态。"],
    ["单像素存储编码", "64 个 signed int8 通道，即 64 B/footprint；官方采用量化 embedding", "FP32；约 4 MiB / patch·月", "AEF 的 8-bit 分发格式是其存储效率关键。玄女需先完成量化精度验证，不能直接照搬。"],
    ["公开 embedding 规模", "每年超过 1.4 万亿 embedding footprint", "320 patch × 6 月 = 1,920 patch-month", "按 1.4 万亿 × 64 B 推算，AEF 基础 embedding payload 至少约 89.6 TB/年；未含 COG 压缩与金字塔差异。"],
    ["训练原始数据存储", "未公开", "原始 127.09 GiB；处理后 37.84 GiB", "AEF 论文没有披露总 TB/PB，不能写成确定值；中国版应以全国试点的每 patch 数据密度估算。"],
    ["发布对象存储总量", "未公开", "海淀生产工作集约 174 GiB", "Google 公开说明使用 COG，但没有披露桶总字节数。中国版应分别预算训练数据、embedding、索引和备份。"],
    ["中国 30% 面积的 10 m 发布换算", "按 AEF 64 B/int8：约 1.84 TB/年；6 个月逐月约 11.06 TB", "当前 P10C FP32：6 个月约 44 TB；若验证 FP16，约 22 TB", "量化是明确的工程优化方向；需要独立评估下游任务精度、检索和变化检测。"],
]


def set_cell_text(cell, text: str, bold: bool = False) -> None:
    cell.text = ""
    for index, line in enumerate(text.split("\n")):
        paragraph = cell.paragraphs[0] if index == 0 else cell.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(0)
        run = paragraph.add_run(line)
        run.bold = bold
        run.font.name = "Arial"
        run.font.size = Pt(8.5)


def add_table(document: Document, title: str, headers: list[str], rows: list[list[str]]) -> None:
    heading = document.add_paragraph()
    heading.paragraph_format.space_before = Pt(8)
    run = heading.add_run(title)
    run.bold = True
    run.font.name = "Arial"
    run.font.size = Pt(11)
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.autofit = True
    for index, header in enumerate(headers):
        set_cell_text(table.rows[0].cells[index], header, bold=True)
    for row in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row):
            set_cell_text(cells[index], value)


def html_table(title: str, headers: list[str], rows: list[list[str]]) -> str:
    def esc(value: str) -> str:
        return (value.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
                .replace("\n", "<br>"))
    head = "".join(f"<th>{esc(h)}</th>" for h in headers)
    body = "".join("<tr>" + "".join(f"<td>{esc(v)}</td>" for v in row) + "</tr>" for row in rows)
    return f"<h2>{title}</h2><table><thead><tr>{head}</tr></thead><tbody>{body}</tbody></table>"


def main() -> None:
    ROOT.mkdir(parents=True, exist_ok=True)
    document = Document()
    section = document.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = section.page_height, section.page_width
    section.top_margin = Cm(1.2)
    section.bottom_margin = Cm(1.2)
    section.left_margin = Cm(1.0)
    section.right_margin = Cm(1.0)

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("中国版 AlphaEarth 方案：可直接复制到 Word 的数据与规模参考表")
    run.bold = True
    run.font.name = "Arial"
    run.font.size = Pt(14)
    note = document.add_paragraph("统计日期：2026-07-16。AEF 未公开的数据容量均明确标注为“未公开”；推算值已在单元格中写明公式与前提。")
    note.runs[0].font.size = Pt(9)
    add_table(document, "表 1  海淀区数据量审计与全国 30% 容量估算", AUDIT_HEADERS, AUDIT_ROWS)
    add_table(document, "表 2  AlphaEarth Foundations 与玄女海淀 P10C 的公开规模参考", AEF_HEADERS, AEF_ROWS)
    refs = document.add_paragraph("来源：AEF 论文 https://arxiv.org/abs/2507.22291；Google DeepMind 介绍 https://deepmind.google/blog/alphaearth-foundations-helps-map-our-planet-in-unprecedented-detail/；Google GCS 数据格式说明 https://developers.google.com/earth-engine/guides/aef_on_gcs_readme")
    refs.runs[0].font.size = Pt(8)
    document.save(DOCX_OUT)

    html = """<!doctype html><html><head><meta charset=\"utf-8\"><title>AlphaEarth Word Copy Tables</title><style>
body { font-family: Arial, 'Microsoft YaHei', sans-serif; margin: 24px; color: #111; }
h1 { font-size: 20px; } h2 { margin-top: 28px; font-size: 16px; }
table { border-collapse: collapse; width: 100%; font-size: 12px; line-height: 1.35; }
th, td { border: 1px solid #333; padding: 7px; vertical-align: top; text-align: left; }
th { background: #d9eaf7; font-weight: 700; }
tr:nth-child(even) td { background: #f7fbff; }
.note { font-size: 12px; } </style></head><body>
<h1>中国版 AlphaEarth 方案：可直接复制到 Word 的数据与规模参考表</h1>
<p class=\"note\">统计日期：2026-07-16。AEF 未公开的数据容量均明确标注为“未公开”；推算值已在单元格中写明公式与前提。</p>
""" + html_table("表 1  海淀区数据量审计与全国 30% 容量估算", AUDIT_HEADERS, AUDIT_ROWS) + html_table("表 2  AlphaEarth Foundations 与玄女海淀 P10C 的公开规模参考", AEF_HEADERS, AEF_ROWS) + """<p class=\"note\">来源：AEF 论文 https://arxiv.org/abs/2507.22291；Google DeepMind 介绍 https://deepmind.google/blog/alphaearth-foundations-helps-map-our-planet-in-unprecedented-detail/；Google GCS 数据格式说明 https://developers.google.com/earth-engine/guides/aef_on_gcs_readme</p></body></html>"""
    HTML_OUT.write_text(html, encoding="utf-8")
    print(DOCX_OUT)
    print(HTML_OUT)


if __name__ == "__main__":
    main()
