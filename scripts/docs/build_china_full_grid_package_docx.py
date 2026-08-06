#!/usr/bin/env python3
"""Build the Chinese README DOCX for the China full 1,280 m grid package."""

from __future__ import annotations

import argparse
import json
from pathlib import Path
from typing import Any, Mapping

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

ROOT = Path(__file__).resolve().parents[2]
DEFAULT_PRIOR_SAMPLING_SOURCE = ROOT / "docs/data/先读我_中国季度嵌入采样说明_20260727.md"
BLACK = RGBColor(0, 0, 0)


def _set_run_font(run: Any, size: float, *, bold: bool = False) -> None:
    """Apply the approved bilingual Word typography to one run."""
    run.font.name = "Times New Roman"
    fonts = run._element.get_or_add_rPr().get_or_add_rFonts()
    fonts.set(qn("w:ascii"), "Times New Roman")
    fonts.set(qn("w:hAnsi"), "Times New Roman")
    fonts.set(qn("w:cs"), "Times New Roman")
    fonts.set(qn("w:eastAsia"), "SimSun")
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = BLACK


def _configure_document(document: Document) -> None:
    normal = document.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "SimSun")
    normal.font.size = Pt(12)
    normal.font.color.rgb = BLACK
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.space_after = Pt(6)

    for style_name, size in (("Title", 16), ("Heading 1", 14), ("Heading 2", 14)):
        style = document.styles[style_name]
        style.font.name = "Times New Roman"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "SimSun")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = BLACK

    document.core_properties.title = "中国 1,280 m 全国完整父网格包"
    document.core_properties.subject = "全国完整父网格、采样成员和空间审计说明"
    document.core_properties.author = "玄女 Embedding 项目组"


def _add_text(document: Document, text: str, *, style: str | None = None) -> None:
    paragraph = document.add_paragraph(style=style)
    _set_run_font(paragraph.add_run(text), 14 if style else 12, bold=style is not None)


def _prior_sampling_excerpt(path: Path) -> str:
    if not path.exists():
        raise FileNotFoundError(f"prior sampling explanation does not exist: {path}")
    for line in path.read_text(encoding="utf-8").splitlines():
        candidate = line.strip()
        if candidate and not candidate.startswith(("#", "|", "!", "-", "版本：")):
            return candidate
    raise ValueError(f"prior sampling explanation has no prose: {path}")


def build_package_docx(
    output_path: str | Path,
    *,
    package_metadata: Mapping[str, Any],
    prior_sampling_source: str | Path = DEFAULT_PRIOR_SAMPLING_SOURCE,
) -> Path:
    """Write the combined package explanation with the approved Word typography."""
    output_path = Path(output_path)
    all_count = int(package_metadata["all_count"])
    sampled_count = int(package_metadata["sampled_count"])
    unsampled_count = int(package_metadata["unsampled_count"])
    if all_count != sampled_count + unsampled_count:
        raise ValueError(
            "package metadata must satisfy all_count = sampled_count + unsampled_count"
        )
    audit = package_metadata.get("membership_audit", {})
    prior_excerpt = _prior_sampling_excerpt(Path(prior_sampling_source))

    document = Document()
    _configure_document(document)
    title = document.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_run_font(title.add_run("先读我：中国 1,280 m 全国完整父网格包"), 16, bold=True)

    _add_text(document, "一、这个包是什么", style="Heading 1")
    _add_text(
        document,
        f"本包交付全国完整父网格，共 {all_count:,} 个 1,280 m × 1,280 m parent patch。"
        "它是长期、可复现的空间索引，不包含遥感影像、季度质量结论或 embedding。",
    )

    _add_text(document, "二、与既有采样候选的关系", style="Heading 1")
    _add_text(document, f"既有采样说明指出：{prior_excerpt}")
    _add_text(
        document,
        f"其中 {sampled_count:,} 个为已采样 patch，按 grid_epsg + grid_col + grid_row "
        "与父网格精确匹配；它们保留原有采样层、质量复核和训练准入限制。",
    )
    _add_text(
        document,
        f"其余 {unsampled_count:,} 个为未采样 patch，只是完整父网格的其余成员，"
        "不是负样本，也不代表已经下载、验收或可直接用于训练。",
    )
    _add_text(
        document,
        "成员关系必须满足 all = sampled + unsampled，且两类 patch 没有共享的 parent_key。",
    )

    _add_text(document, "三、坐标系和网格规则", style="Heading 1")
    _add_text(
        document,
        "网格在 WGS84 / UTM 北半球米制坐标系中建立，覆盖 EPSG:32643 至 EPSG:32653；"
        "例如 UTM 43N 为 EPSG:32643。每个 patch 以所属 UTM 的 (0 m, 0 m) 为原点对齐。",
    )
    _add_text(
        document,
        "交付图层统一为 WGS84 经纬度 EPSG:4326，坐标轴顺序为 X=经度、Y=纬度。"
        "EPSG:4326 适合显示和交换；计算长度、面积、缓冲和 overlap 时须转回 grid_epsg。",
    )

    _add_text(document, "四、文件与预览", style="Heading 1")
    _add_text(
        document,
        "all/、sampled/ 和 unsampled/ 保存分区 GeoParquet 与分片 Shapefile；"
        "china_full_1280m_macrocell_index.gpkg 含 macrocells 和 utm_zones 两个索引图层。",
    )
    _add_text(
        document,
        "national_patch_density_by_utm_zone.png 检查各 UTM 分区密度；"
        "local_1280m_grid_sampled_unsampled.png 同时显示两类局部 patch；"
        "utm_owner_zone_seams.png 说明 owner-zone 接缝和接受的跨区 overlap 规则。",
    )

    _add_text(document, "五、审计结论和限制", style="Heading 1")
    audit_state = "通过" if bool(audit.get("passed")) else "未通过"
    matched_count = int(audit.get("matched", 0))
    _add_text(
        document,
        f"本次成员审计状态：{audit_state}；精确匹配的采样成员数为 {matched_count:,}。"
        "宏网格计数之和必须等于完整 GeoParquet 的 all 计数。",
    )
    _add_text(
        document,
        "边界和 UTM 接缝附近的 footprint 可以跨出 owner-zone 边界。验收要求重叠只发生"
        "在相邻 UTM owner zone 的边界窄带、两侧中心归属正确、非相邻重叠为零，并且全国"
        "累计重复覆盖面积不超过父网格总面积的 0.1%。重新采样、下载和训练准入仍须遵循"
        "既有采样说明及相应质量审计。",
    )

    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(output_path)
    return output_path


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--metadata", type=Path, required=True)
    parser.add_argument("--prior-sampling-source", type=Path, default=DEFAULT_PRIOR_SAMPLING_SOURCE)
    args = parser.parse_args()
    metadata = json.loads(args.metadata.read_text(encoding="utf-8"))
    build_package_docx(
        args.output,
        package_metadata=metadata,
        prior_sampling_source=args.prior_sampling_source,
    )
    print(args.output)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
