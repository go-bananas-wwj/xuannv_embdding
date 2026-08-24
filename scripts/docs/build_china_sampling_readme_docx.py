#!/usr/bin/env python3
"""Build the concise China sampling README DOCX and its Chinese preview map."""

from __future__ import annotations

import re
import shutil
from pathlib import Path

import geopandas as gpd
import matplotlib
import pandas as pd

matplotlib.use("Agg")
import matplotlib.pyplot as plt
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor
from matplotlib import font_manager

ROOT = Path(__file__).resolve().parents[2]
SOURCE = ROOT / "docs/data/先读我_中国季度嵌入采样说明_20260727.md"
ASSET_DIR = ROOT / "docs/plans/assets/china_quarterly_sampling_20260727"
MAP = ASSET_DIR / "china_quarterly_62000_candidate_preview_zh.png"
FLOW = ASSET_DIR / "sampling_composition_and_workflow_readme.png"
REGISTRY = Path(
    "/data/xuannv_embedding/outputs/"
    "china_quarterly_sampling_62000_candidate_20260726/"
    "china_quarterly_62000_candidate_registry.csv"
)
COUNTRY = Path(
    "/data/xuannv_embedding/raw/china_v1/geoboundaries/"
    "geoBoundaries-CHN-ADM0.geojson"
)
ADM1 = Path(
    "/data/xuannv_embedding/raw/china_v1/geoboundaries/"
    "geoBoundaries-CHN-ADM1.geojson"
)
OUTPUT = ROOT / "docs/plans/word/先读我.docx"
DELIVERY = Path(
    "/data/xuannv_embedding/outputs/"
    "china_quarterly_sampling_62000_candidate_20260726/先读我.docx"
)
SONG_FONT = "/usr/share/fonts/google-noto-cjk/NotoSerifCJK-Regular.ttc"


def set_run_font(run, size: float = 12, bold: bool | None = None) -> None:
    run.font.name = "Times New Roman"
    fonts = run._element.get_or_add_rPr().get_or_add_rFonts()
    for key in ("ascii", "hAnsi", "cs"):
        fonts.set(qn(f"w:{key}"), "Times New Roman")
    fonts.set(qn("w:eastAsia"), "宋体")
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor(0, 0, 0)
    if bold is not None:
        run.bold = bold


def set_cell_shading(cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), fill)


def add_page_number(section) -> None:
    paragraph = section.footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(paragraph.add_run("第 "), 9)
    run = OxmlElement("w:r")
    for kind, text in (
        ("begin", None),
        (None, " PAGE "),
        ("separate", None),
        (None, "1"),
        ("end", None),
    ):
        if kind:
            node = OxmlElement("w:fldChar")
            node.set(qn("w:fldCharType"), kind)
        else:
            node = OxmlElement("w:instrText" if text == " PAGE " else "w:t")
            node.text = text
        run.append(node)
    paragraph._p.append(run)
    set_run_font(paragraph.add_run(" 页"), 9)


def configure(document: Document) -> None:
    section = document.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.5)
    add_page_number(section)

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(header.add_run("中国版季度嵌入采样数据说明"), 9)

    normal = document.styles["Normal"]
    normal.font.name = "Times New Roman"
    for key in ("ascii", "hAnsi"):
        normal._element.rPr.rFonts.set(qn(f"w:{key}"), "Times New Roman")
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    normal.font.size = Pt(12)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.space_after = Pt(6)

    for name, size, before, after in (
        ("Title", 16, 0, 14),
        ("Heading 1", 14, 14, 7),
        ("Heading 2", 14, 10, 5),
    ):
        style = document.styles[name]
        style.font.name = "Times New Roman"
        for key in ("ascii", "hAnsi"):
            style._element.rPr.rFonts.set(qn(f"w:{key}"), "Times New Roman")
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)

    document.core_properties.title = "先读我：中国版季度嵌入采样数据说明"
    document.core_properties.author = "玄女 Embedding 项目组"


def add_inline(paragraph, text: str, size: float = 12) -> None:
    pattern = re.compile(r"(\*\*.+?\*\*|`.+?`)")
    cursor = 0
    for match in pattern.finditer(text):
        if match.start() > cursor:
            set_run_font(paragraph.add_run(text[cursor : match.start()]), size)
        token = match.group(0)
        run = paragraph.add_run(token[2:-2] if token.startswith("**") else token[1:-1])
        set_run_font(run, size, token.startswith("**"))
        cursor = match.end()
    if cursor < len(text):
        set_run_font(paragraph.add_run(text[cursor:]), size)


def join_wrapped_lines(parts: list[str]) -> str:
    """Join Markdown source wraps without inserting spaces between Chinese text."""
    def cjk_or_punctuation(character: str) -> bool:
        return (
            "\u2e80" <= character <= "\u9fff"
            or "\u3000" <= character <= "\u303f"
            or "\uff00" <= character <= "\uffef"
        )

    result = parts[0]
    for part in parts[1:]:
        needs_space = bool(
            result
            and part
            and not (
                cjk_or_punctuation(result[-1])
                and cjk_or_punctuation(part[0])
            )
        )
        result += (" " if needs_space else "") + part
    return result


def add_table(document: Document, rows: list[list[str]]) -> None:
    columns = max(len(row) for row in rows)
    table = document.add_table(rows=len(rows), cols=columns)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    for row_index, row in enumerate(rows):
        tr_pr = table.rows[row_index]._tr.get_or_add_trPr()
        tr_pr.append(OxmlElement("w:cantSplit"))
        if row_index == 0:
            header = OxmlElement("w:tblHeader")
            header.set(qn("w:val"), "true")
            tr_pr.append(header)
        for column_index in range(columns):
            cell = table.cell(row_index, column_index)
            cell.text = ""
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if row_index == 0:
                set_cell_shading(cell, "E7E6E3")
            paragraph = cell.paragraphs[0]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            value = row[column_index] if column_index < len(row) else ""
            add_inline(paragraph, value, 12)
            if row_index == 0:
                for run in paragraph.runs:
                    run.bold = True


def render_map() -> None:
    ASSET_DIR.mkdir(parents=True, exist_ok=True)
    frame = pd.read_csv(REGISTRY)
    country = gpd.read_file(COUNTRY).to_crs("EPSG:4326")
    adm1 = gpd.read_file(ADM1).to_crs("EPSG:4326")
    font_manager.fontManager.addfont(SONG_FONT)
    chinese = font_manager.FontProperties(fname=SONG_FONT)
    colors = {
        "base_expected_1pct": ("#B44A3D", "基础1%空间样本（57,405）", 1.0),
        "utm_spatial_balance_supplement": ("#3979B6", "UTM空间均衡（2,595）", 3.0),
        "coastal_supplement": ("#7B4A9E", "海岸线（500）", 4.0),
        "semantic_and_difficult_supplement": (
            "#3C8C62",
            "语义与困难区域候选（1,500）",
            3.5,
        ),
    }
    figure, axis = plt.subplots(figsize=(12.2, 8.0), dpi=220, facecolor="white")
    country.plot(ax=axis, color="#F5F5F3", edgecolor="#333333", linewidth=0.6, zorder=1)
    adm1.boundary.plot(ax=axis, color="#B5B5B5", linewidth=0.18, zorder=2)
    for layer, (color, label, size) in colors.items():
        part = frame.loc[frame["sampling_layer"] == layer]
        axis.scatter(
            part["longitude"],
            part["latitude"],
            s=size,
            c=color,
            alpha=0.72,
            linewidths=0,
            label=label,
            zorder=3,
        )
    min_x, min_y, max_x, max_y = country.total_bounds
    axis.set_xlim(min_x - 1.2, max_x + 1.2)
    axis.set_ylim(min_y - 1.2, max_y + 1.2)
    axis.set_aspect("equal")
    axis.set_axis_off()
    axis.set_title("全国 62,000 个季度嵌入采样候选", fontproperties=chinese, fontsize=16)
    legend = axis.legend(
        loc="upper left",
        frameon=True,
        prop=chinese,
        fontsize=9,
        markerscale=2.6,
    )
    legend.get_frame().set_edgecolor("#888888")
    axis.text(
        0.01,
        0.015,
        "内部技术预览：候选位置已通过几何审计，季度影像质量复核尚未完成",
        transform=axis.transAxes,
        fontsize=8.5,
        color="#303030",
        fontproperties=chinese,
        bbox={"facecolor": "white", "edgecolor": "#888888", "pad": 4.0},
    )
    figure.tight_layout()
    figure.savefig(MAP, bbox_inches="tight", facecolor="white")
    plt.close(figure)


def render_flow() -> None:
    ASSET_DIR.mkdir(parents=True, exist_ok=True)
    font_manager.fontManager.addfont(SONG_FONT)
    chinese = font_manager.FontProperties(fname=SONG_FONT)
    labels = ["基础1%空间样本", "UTM空间均衡", "海岸线", "语义与困难候选"]
    values = [57405, 2595, 500, 1500]
    colors = ["#B44A3D", "#3979B6", "#7B4A9E", "#3C8C62"]
    figure = plt.figure(figsize=(12.2, 6.4), dpi=220, facecolor="white")
    grid = figure.add_gridspec(2, 1, height_ratios=[1.0, 1.2], hspace=0.48)
    axis = figure.add_subplot(grid[0])
    left = 0
    annotation_y = [0.0, 0.66, 1.12, 0.30]
    annotation_x = [0, 57500, 60050, 61800]
    annotation_align = ["center", "center", "center", "right"]
    for index, (label, value, color) in enumerate(zip(labels, values, colors, strict=True)):
        axis.barh([0], [value], left=[left], color=color, height=0.5, edgecolor="white")
        if index == 0:
            axis.text(
                left + value / 2,
                0,
                f"{label}  {value:,}",
                color="white",
                ha="center",
                va="center",
                fontsize=10,
                fontproperties=chinese,
            )
        else:
            axis.annotate(
                f"{label}  {value:,}",
                xy=(left + value / 2, 0.25),
                xytext=(annotation_x[index], annotation_y[index]),
                ha=annotation_align[index],
                va="bottom",
                fontsize=8.3,
                fontproperties=chinese,
                arrowprops={"arrowstyle": "-", "color": "#555555", "lw": 0.8},
            )
        left += value
    axis.set_xlim(0, 62000)
    axis.set_ylim(-0.55, 1.42)
    axis.set_yticks([])
    axis.set_xticks([0, 10000, 20000, 30000, 40000, 50000, 62000])
    axis.set_xticklabels(
        ["0", "1万", "2万", "3万", "4万", "5万", "6.2万"],
        fontproperties=chinese,
    )
    axis.set_title("62,000 个候选 patch 的组成", fontproperties=chinese, fontsize=15)
    for side in ("top", "right", "left"):
        axis.spines[side].set_visible(False)

    flow = figure.add_subplot(grid[1])
    flow.axis("off")
    steps = [
        ("历史空间候选骨架", "修复UTM与重复"),
        ("基础空间层", "57,405"),
        ("空间均衡补样", "+2,595"),
        ("海岸候选", "+500"),
        ("语义困难候选", "+1,500"),
        ("季度影像质量审计", "尚未完成"),
        ("冻结训练清单", "目标62,000"),
    ]
    x_positions = [0.07, 0.215, 0.36, 0.505, 0.65, 0.80, 0.94]
    for index, ((title, subtitle), x) in enumerate(zip(steps, x_positions, strict=True)):
        flow.text(
            x,
            0.55,
            f"{title}\n{subtitle}",
            ha="center",
            va="center",
            fontsize=8.0,
            fontproperties=chinese,
            bbox={
                "boxstyle": "round,pad=0.5,rounding_size=0.06",
                "facecolor": "#F5F5F3",
                "edgecolor": "#555555",
                "linewidth": 0.8,
            },
        )
        if index < len(steps) - 1:
            flow.annotate(
                "",
                xy=(x_positions[index + 1] - 0.06, 0.55),
                xytext=(x + 0.06, 0.55),
                arrowprops={"arrowstyle": "->", "lw": 1.0, "color": "#666666"},
            )
    flow.text(
        0.5,
        0.1,
        "当前交付到候选注册表阶段；完成八季度影像质量复核后才能用于正式训练",
        ha="center",
        va="center",
        fontsize=9.5,
        fontproperties=chinese,
    )
    figure.savefig(FLOW, bbox_inches="tight", facecolor="white")
    plt.close(figure)


def build_docx() -> None:
    document = Document()
    configure(document)
    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    index = 0
    while index < len(lines):
        stripped = lines[index].strip()
        if not stripped:
            index += 1
            continue
        heading = re.match(r"^(#{1,3})\s+(.+)$", stripped)
        if heading:
            level = len(heading.group(1))
            style = "Title" if level == 1 else "Heading 1" if level == 2 else "Heading 2"
            paragraph = document.add_paragraph(style=style)
            paragraph.alignment = (
                WD_ALIGN_PARAGRAPH.CENTER if level == 1 else WD_ALIGN_PARAGRAPH.LEFT
            )
            size = 16 if level == 1 else 14
            set_run_font(paragraph.add_run(heading.group(2)), size, True)
            index += 1
            continue
        if stripped.startswith("|") and stripped.endswith("|"):
            table_lines = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                table_lines.append(lines[index].strip())
                index += 1
            rows = [
                [cell.strip() for cell in line.strip("|").split("|")]
                for line in table_lines
            ]
            if len(rows) > 1 and all(re.fullmatch(r":?-{3,}:?", c) for c in rows[1]):
                rows.pop(1)
            add_table(document, rows)
            continue
        image = re.fullmatch(r"!\[([^\]]*)\]\(([^)]+)\)", stripped)
        if image:
            path = (SOURCE.parent / image.group(2)).resolve()
            paragraph = document.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.add_run().add_picture(str(path), width=Inches(6.1))
            index += 1
            continue
        numbered = re.match(r"^(\d+)(?:\.\s+|、\s*)(.+)$", stripped)
        if numbered:
            list_lines = [numbered.group(2)]
            index += 1
            while index < len(lines):
                candidate = lines[index].strip()
                if (
                    not candidate
                    or candidate.startswith(("#", "|", "!["))
                    or re.match(r"^\d+(?:\.\s+|、\s*)", candidate)
                ):
                    break
                list_lines.append(candidate)
                index += 1
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.left_indent = Cm(0.6)
            add_inline(paragraph, f"{numbered.group(1)}. {join_wrapped_lines(list_lines)}")
            continue
        paragraph_lines = [stripped]
        index += 1
        while index < len(lines):
            candidate = lines[index].strip()
            if (
                not candidate
                or candidate.startswith(("#", "|", "!["))
                or re.match(r"^\d+(?:\.\s+|、\s*)", candidate)
            ):
                break
            paragraph_lines.append(candidate)
            index += 1
        paragraph = document.add_paragraph()
        if paragraph_lines[0].startswith("**图"):
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_inline(paragraph, join_wrapped_lines(paragraph_lines))

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT)
    shutil.copy2(OUTPUT, DELIVERY)


if __name__ == "__main__":
    render_map()
    render_flow()
    build_docx()
    print(OUTPUT)
    print(DELIVERY)
