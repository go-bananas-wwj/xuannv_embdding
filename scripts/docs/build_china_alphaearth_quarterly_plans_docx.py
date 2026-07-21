#!/usr/bin/env python3
"""Build the two China AlphaEarth planning DOCX files from Markdown sources."""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[2]
SOURCES = [
    ROOT / "docs/plans/china_alphaearth_existing_resources_quarterly_2020_2021_20260721_zh.md",
    ROOT / "docs/plans/china_alphaearth_high_resolution_quarterly_2020_2021_20260721_zh.md",
]
OUTPUT_DIR = ROOT / "docs/plans/word"
BLUE = "1F4E78"
LIGHT_BLUE = "D9EAF7"
LIGHT_GRAY = "F2F5F7"
TEXT = RGBColor(31, 41, 55)


def set_run_font(run, name="Microsoft YaHei", size=None, bold=None, color=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = color


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_text(cell, text, header=False):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if header else WD_ALIGN_PARAGRAPH.LEFT
    add_inline(paragraph, text)
    for run in paragraph.runs:
        set_run_font(
            run,
            size=9,
            bold=True if header else run.bold,
            color=RGBColor(255, 255, 255) if header else TEXT,
        )
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_hyperlink(paragraph, label, url):
    relation_id = paragraph.part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relation_id)
    run = OxmlElement("w:r")
    run_props = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    run_props.extend([color, underline])
    text = OxmlElement("w:t")
    text.text = label
    run.extend([run_props, text])
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


INLINE_RE = re.compile(r"(\*\*.+?\*\*|`.+?`|\[[^\]]+\]\([^)]+\))")


def add_inline(paragraph, text):
    cursor = 0
    for match in INLINE_RE.finditer(text):
        if match.start() > cursor:
            run = paragraph.add_run(text[cursor : match.start()])
            set_run_font(run, size=10.5, color=TEXT)
        token = match.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            set_run_font(run, size=10.5, bold=True, color=TEXT)
        elif token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            set_run_font(run, name="Consolas", size=9.5, color=TEXT)
        else:
            label, url = re.match(r"\[([^\]]+)\]\(([^)]+)\)", token).groups()
            add_hyperlink(paragraph, label, url)
        cursor = match.end()
    if cursor < len(text):
        run = paragraph.add_run(text[cursor:])
        set_run_font(run, size=10.5, color=TEXT)


def add_page_number(section):
    paragraph = section.footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("第 ")
    set_run_font(run, size=9, color=RGBColor(100, 116, 139))
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    paragraph._p.append(field)
    run = paragraph.add_run(" 页")
    set_run_font(run, size=9, color=RGBColor(100, 116, 139))


def configure_document(document, title):
    section = document.sections[0]
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(1.8)
    section.left_margin = Cm(2.1)
    section.right_margin = Cm(2.1)
    section.header_distance = Cm(0.8)
    section.footer_distance = Cm(0.8)
    add_page_number(section)

    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = TEXT
    normal.paragraph_format.line_spacing = 1.35
    normal.paragraph_format.space_after = Pt(5)

    heading_specs = {
        "Title": (22, RGBColor(20, 54, 86), 18, 12),
        "Heading 1": (15, RGBColor(31, 78, 120), 14, 6),
        "Heading 2": (12, RGBColor(46, 95, 130), 10, 4),
        "Heading 3": (11, RGBColor(55, 65, 81), 8, 3),
    }
    for style_name, (size, color, before, after) in heading_specs.items():
        style = styles[style_name]
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = color
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)

    document.core_properties.title = title
    document.core_properties.subject = "中国版 Alpha Earth 2020--2021 季度 embedding 技术方案"
    document.core_properties.author = "玄女 Embedding 项目组"


def parse_table(lines):
    rows = []
    for line in lines:
        rows.append([cell.strip() for cell in line.strip().strip("|").split("|")])
    if len(rows) >= 2 and all(re.fullmatch(r":?-{3,}:?", cell) for cell in rows[1]):
        rows.pop(1)
    return rows


def add_table(document, rows):
    column_count = max(len(row) for row in rows)
    table = document.add_table(rows=len(rows), cols=column_count)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    for row_index, row in enumerate(rows):
        for column_index in range(column_count):
            value = row[column_index] if column_index < len(row) else ""
            cell = table.cell(row_index, column_index)
            set_cell_text(cell, value, header=row_index == 0)
            if row_index == 0:
                set_cell_shading(cell, BLUE)
            elif row_index % 2 == 0:
                set_cell_shading(cell, LIGHT_GRAY)
    document.add_paragraph().paragraph_format.space_after = Pt(0)


def build_docx(source):
    document = Document()
    lines = source.read_text(encoding="utf-8").splitlines()
    title = lines[0].lstrip("# ")
    configure_document(document, title)

    index = 0
    in_code = False
    code_lines = []
    while index < len(lines):
        line = lines[index]
        stripped = line.strip()
        if stripped.startswith("```"):
            if in_code:
                paragraph = document.add_paragraph()
                paragraph.paragraph_format.left_indent = Cm(0.8)
                paragraph.paragraph_format.right_indent = Cm(0.8)
                paragraph.paragraph_format.space_before = Pt(4)
                paragraph.paragraph_format.space_after = Pt(7)
                run = paragraph.add_run("\n".join(code_lines))
                set_run_font(run, name="Consolas", size=9, color=RGBColor(51, 65, 85))
                in_code = False
                code_lines = []
            else:
                in_code = True
            index += 1
            continue
        if in_code:
            code_lines.append(line)
            index += 1
            continue
        if not stripped:
            index += 1
            continue
        if stripped.startswith("|") and stripped.endswith("|"):
            table_lines = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                table_lines.append(lines[index])
                index += 1
            add_table(document, parse_table(table_lines))
            continue
        image_match = re.fullmatch(r"!\[([^\]]*)\]\(([^)]+)\)", stripped)
        if image_match:
            image_path = (source.parent / image_match.group(2)).resolve()
            paragraph = document.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.add_run().add_picture(str(image_path), width=Inches(6.2))
            index += 1
            continue
        heading = re.match(r"^(#{1,3})\s+(.+)$", stripped)
        if heading:
            level = len(heading.group(1))
            style = "Title" if level == 1 else f"Heading {level - 1}"
            paragraph = document.add_paragraph(style=style)
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if level == 1 else WD_ALIGN_PARAGRAPH.LEFT
            add_inline(paragraph, heading.group(2))
            index += 1
            continue
        if stripped.startswith(">"):
            quote_lines = []
            while index < len(lines) and lines[index].strip().startswith(">"):
                quote_lines.append(lines[index].strip().lstrip("> "))
                index += 1
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.left_indent = Cm(0.7)
            paragraph.paragraph_format.right_indent = Cm(0.7)
            paragraph.paragraph_format.space_before = Pt(6)
            paragraph.paragraph_format.space_after = Pt(8)
            add_inline(paragraph, " ".join(quote_lines))
            p_pr = paragraph._p.get_or_add_pPr()
            shading = OxmlElement("w:shd")
            shading.set(qn("w:fill"), LIGHT_BLUE)
            p_pr.append(shading)
            continue
        bullet = re.match(r"^[-*]\s+(.+)$", stripped)
        numbered = re.match(r"^\d+\.\s+(.+)$", stripped)
        if bullet or numbered:
            paragraph = document.add_paragraph(
                style="List Bullet" if bullet else "List Number"
            )
            add_inline(paragraph, (bullet or numbered).group(1))
            index += 1
            continue

        paragraph_lines = [stripped]
        index += 1
        while index < len(lines):
            candidate = lines[index].strip()
            if (
                not candidate
                or candidate.startswith(("#", "|", ">", "```", "![", "- ", "* "))
                or re.match(r"^\d+\.\s+", candidate)
            ):
                break
            paragraph_lines.append(candidate)
            index += 1
        paragraph = document.add_paragraph()
        if paragraph_lines[0].startswith("**图"):
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_inline(paragraph, " ".join(paragraph_lines))

    output = OUTPUT_DIR / f"{source.stem}.docx"
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    document.save(output)
    return output


def main():
    for source in SOURCES:
        output = build_docx(source)
        print(output)


if __name__ == "__main__":
    main()
