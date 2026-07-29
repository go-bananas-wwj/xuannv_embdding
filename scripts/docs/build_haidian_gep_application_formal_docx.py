from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


ROOT = Path(__file__).resolve().parents[2]
SOURCE = ROOT / "docs/projects/haidian_gep_xuannv_application_formal_20260729_zh.md"
OUTPUT = ROOT / "docs/projects/haidian_gep_xuannv_application_formal_20260729_zh.docx"


def set_run_font(run, size: Pt, bold: bool = False) -> None:
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    run.font.size = size
    run.font.bold = bold
    run.font.color.rgb = None


def add_text(paragraph, text: str, size: Pt, bold: bool = False) -> None:
    run = paragraph.add_run(text)
    set_run_font(run, size, bold=bold)


def add_route_table(document: Document) -> None:
    table = document.add_table(rows=1, cols=5)
    table.autofit = False
    labels = [
        "多源遥感与\n生态辅助数据",
        "数据对齐与\n质量管理",
        "月度地表状态\n嵌入与分类",
        "三类服务实物量\n与价值量核算",
        "试点报告、专题图\n与核算辅助工具",
    ]
    for idx, label in enumerate(labels):
        cell = table.cell(0, idx)
        cell.width = Cm(3.05)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        for part_idx, line in enumerate(label.split("\n")):
            if part_idx:
                p.add_run().add_break()
            add_text(p, line, Pt(10.5), bold=True)
        tc_pr = cell._tc.get_or_add_tcPr()
        shading = OxmlElement("w:shd")
        shading.set(qn("w:fill"), "F2F2F2")
        tc_pr.append(shading)
    document.add_paragraph("技术路线：多源数据统一处理，形成月度生态状态基础产品；以生态斑块、子汇水区和绿地影响范围为过程单元，完成三类调节服务价值试点。")


def build_document() -> None:
    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    document = Document()
    section = document.sections[0]
    section.top_margin = Cm(2.54)
    section.bottom_margin = Cm(2.54)
    section.left_margin = Cm(2.7)
    section.right_margin = Cm(2.7)
    normal = document.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    normal.font.size = Pt(12)
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.first_line_indent = Cm(0.74)

    title = lines[0].removeprefix("# ").strip()
    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(36)
    add_text(p, title, Pt(16), bold=True)
    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(36)
    add_text(p, "自立科研项目申报书（技术内容稿）", Pt(14), bold=True)

    for raw in lines[1:]:
        text = raw.strip()
        if not text:
            continue
        if text == "[[TECH_ROUTE]]":
            add_route_table(document)
            continue
        if text.startswith("## "):
            p = document.add_paragraph()
            p.paragraph_format.space_before = Pt(15)
            p.paragraph_format.space_after = Pt(8)
            add_text(p, text[3:], Pt(14), bold=True)
            continue
        if text.startswith("### "):
            p = document.add_paragraph()
            p.paragraph_format.space_before = Pt(9)
            p.paragraph_format.space_after = Pt(4)
            add_text(p, text[4:], Pt(12), bold=True)
            continue
        p = document.add_paragraph()
        p.paragraph_format.space_after = Pt(5)
        p.paragraph_format.line_spacing = 1.5
        p.paragraph_format.first_line_indent = Cm(0.74)
        add_text(p, text, Pt(12))

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_text(footer, "海淀区生态系统调节服务价值核算技术试点研究", Pt(9))
    document.save(OUTPUT)


if __name__ == "__main__":
    build_document()
