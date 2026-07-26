#!/usr/bin/env python3
"""Build editable native PowerPoint slides 4 through 7 for the Pujiang deck."""

from __future__ import annotations

import argparse
from io import BytesIO
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches, Pt
from pujiang_editable_common import (
    BLUE,
    CYAN,
    DEEP_BLUE,
    INK,
    LINE,
    MUTED,
    PALE_BLUE,
    PALE_GRAY,
    RED,
    WHITE,
    add_header,
    add_picture,
    add_rect,
    add_text,
    new_blank_slide,
)

ROOT = Path(__file__).resolve().parents[2]
PRESENTATION_ROOT = ROOT / "docs/presentations/pujiang_202607"
SLIDE01_ROOT = ROOT / "docs/presentations/assets/pujiang_202607/slide01"
EDITABLE_SOURCE_ROOT = ROOT / "docs/presentations/assets/pujiang_202607/editable_sources"
PLATFORM_SCREENSHOT = EDITABLE_SOURCE_ROOT / "haidian_platform_thum.png"
RECON_EXAMPLE = EDITABLE_SOURCE_ROOT / "patch_000093_recon.png"
PLATFORM_POSTER = ROOT / "docs/bp_deck/assets/video/custom_annotation_demo_poster.png"
WORKFLOW_IMAGE = ROOT / "docs/bp_deck/assets/remote_sensing_team_workflow.png"
OBSERVATION_TIMELINE = (
    ROOT
    / "docs/bp_deck/assets/user_materials/patch_000027_2025-08_vs_2025-09_detail.png"
)
PU_QUERY_ASSET_ROOT = (
    ROOT / "docs/production/assets/haidian_pu_query_3polygon_compare_20260726"
)

PAGE_TITLES = {
    4: "玄女嵌入平台：浏览、分析、训练与生成",
    5: "少量标注，快速形成区域级专题图",
    6: "嵌入底座进入遥感智能体工作流",
    7: "云遮挡或观测缺失时，生成指定时刻的遥感参考影像",
}


def _add_panel_frame(slide, x: float, y: float, w: float, h: float) -> None:
    add_rect(slide, x, y, w, h, WHITE, LINE)


def _add_number_badge(slide, number: int, x: float, y: float) -> None:
    add_rect(slide, x, y, 0.38, 0.38, BLUE, BLUE, MSO_SHAPE.OVAL)
    add_text(slide, str(number), x, y + 0.01, 0.38, 0.28, 13, WHITE, True)


def _set_header_title_size(slide, size: float) -> None:
    """Reduce a long header title without changing the shared deck primitive."""
    for paragraph in slide.shapes[1].text_frame.paragraphs:
        for run in paragraph.runs:
            run.font.size = Pt(size)


def _add_cropped_picture(
    slide,
    source_path: Path,
    crop_box: tuple[int, int, int, int],
    x: float,
    y: float,
    w: float,
    h: float,
):
    """Embed one source crop as its own picture relationship in the PPTX."""
    with Image.open(source_path) as source:
        source_width, source_height = source.size
        left, top, right, bottom = crop_box
        if left < 0 or top < 0 or right > source_width or bottom > source_height:
            raise ValueError(f"Crop {crop_box} exceeds {source_path} dimensions {source.size}")
        crop = source.crop(crop_box).convert("RGB")
    stream = BytesIO()
    crop.save(stream, format="PNG")
    stream.seek(0)
    return slide.shapes.add_picture(stream, Inches(x), Inches(y), Inches(w), Inches(h))


def build_page04(prs: Presentation):
    """Build the platform capability page with independent visual assets."""
    slide = new_blank_slide(prs)
    add_header(
        slide,
        PAGE_TITLES[4],
        "同一份月度嵌入，从离线数据文件变成可浏览、可训练、可调用的业务能力",
    )
    add_text(
        slide,
        "平台主界面与关键入口：嵌入浏览、专题分析、自定义训练、时相生成",
        0.55,
        1.03,
        7.0,
        0.28,
        14,
        MUTED,
        True,
    )

    _add_panel_frame(slide, 0.5, 1.4, 7.1, 4.55)
    add_picture(slide, PLATFORM_SCREENSHOT, 0.65, 1.57, 6.8, 4.2, "contain")
    add_text(slide, "海淀区主界面", 0.7, 5.63, 2.1, 0.2, 13, DEEP_BLUE, True)
    add_rect(slide, 2.0, 2.0, 2.18, 1.48, WHITE, CYAN)
    add_picture(slide, SLIDE01_ROOT / "haidian_pca.png", 2.04, 2.04, 2.1, 1.4, "cover")
    add_text(slide, "月度嵌入 PCA", 2.03, 3.52, 2.1, 0.2, 11, DEEP_BLUE, True)
    add_rect(slide, 4.75, 3.52, 1.68, 1.22, WHITE, CYAN)
    add_picture(slide, PLATFORM_POSTER, 4.79, 3.56, 1.6, 1.14, "cover")
    add_text(slide, "标注演示", 4.78, 4.76, 1.6, 0.2, 11, DEEP_BLUE, True)

    functions = [
        ("嵌入浏览", "按区域、月份、任务和 patch\n查看嵌入与专题结果"),
        ("专题分析", "建筑、道路、水体、变化等\n结果回填地图统一查看"),
        ("自定义训练", "少量标注后创建或加载\n轻量下游模型"),
        ("时相生成", "选择目标时刻，生成缺测或\n云遮挡场景的参考影像"),
    ]
    for index, (title, detail) in enumerate(functions, start=1):
        y = 1.4 + (index - 1) * 1.08
        add_rect(slide, 8.0, y, 4.78, 0.9, PALE_BLUE, LINE, MSO_SHAPE.ROUNDED_RECTANGLE)
        _add_number_badge(slide, index, 8.23, y + 0.26)
        add_text(slide, title, 8.78, y + 0.14, 1.35, 0.22, 16, DEEP_BLUE, True)
        add_text(slide, detail, 10.2, y + 0.12, 2.25, 0.5, 12, MUTED)

    add_rect(slide, 0.5, 6.35, 12.28, 0.55, PALE_GRAY, PALE_GRAY, MSO_SHAPE.ROUNDED_RECTANGLE)
    add_text(
        slide,
        "平台入口：112.111.7.74:1114/embedEarth?regionId=haidian  |  "
        "嵌入浏览 -> 任务训练 -> 全域推理 -> 地图复核",
        0.68,
        6.51,
        11.9,
        0.18,
        12,
        INK,
        True,
    )
    return slide


def build_page05(prs: Presentation):
    """Build the few-shot candidate comparison with 18 independent picture panels."""
    slide = new_blank_slide(prs)
    add_header(
        slide,
        PAGE_TITLES[5],
        "用户每类只圈 3 个目标多边形；玄女、AEF 与传统特征共享同一标注和 PU+Query 流程",
    )
    add_rect(slide, 0.5, 1.08, 12.28, 0.62, PALE_BLUE, LINE, MSO_SHAPE.ROUNDED_RECTANGLE)
    add_text(
        slide,
        "PU：从未标区域中自动寻找可靠背景；Query：在新 patch 中用高置信目标轻量校准。\n"
        "候选图仅做各方法单 patch 像素百分位显示：白低、红高；"
        "色深不可跨方法比较，也不改变汇总指标或二值阈值。",
        0.66,
        1.17,
        11.96,
        0.4,
        12,
        INK,
        True,
    )

    columns = ("3 个圈选", "测试影像", "玄女候选", "AEF 候选", "传统候选", "真实标签")
    image_x = 1.32
    image_width = 1.84
    image_gap = 0.15
    for index, title in enumerate(columns):
        x = image_x + index * (image_width + image_gap)
        add_text(slide, title, x, 1.91, image_width, 0.24, 14, DEEP_BLUE, True)

    crop_left = 112
    crop_top = 40
    crop_width = 232
    crop_height = 142
    crop_gap = 243
    tasks = (("建筑物", "building"), ("道路", "road"), ("水体", "water"))
    for row, (task, filename_stem) in enumerate(tasks):
        y = 2.24 + row * 1.46
        add_rect(slide, 0.45, y, 0.72, 1.12, BLUE, BLUE, MSO_SHAPE.ROUNDED_RECTANGLE)
        add_text(slide, task, 0.52, y + 0.43, 0.58, 0.22, 13, WHITE, True)
        source = PU_QUERY_ASSET_ROOT / f"{filename_stem}_slide_row.png"
        for column in range(6):
            x = image_x + column * (image_width + image_gap)
            add_rect(slide, x - 0.02, y - 0.02, image_width + 0.04, 1.16, WHITE, LINE)
            crop = (
                crop_left + column * crop_gap,
                crop_top,
                crop_left + crop_width + column * crop_gap,
                crop_top + crop_height,
            )
            _add_cropped_picture(slide, source, crop, x, y, image_width, 1.12)

    add_rect(slide, 0.5, 6.72, 12.28, 0.48, PALE_GRAY, PALE_GRAY, MSO_SHAPE.ROUNDED_RECTANGLE)
    add_text(
        slide,
        "三行依次为建筑物、道路、水体；本页显示候选排序，第 8 页报告独立测试集二值指标。",
        0.72,
        6.86,
        11.85,
        0.22,
        10,
        MUTED,
        True,
    )
    return slide


def build_page06(prs: Presentation):
    """Build the agent workflow page with one image and editable workflow cards."""
    slide = new_blank_slide(prs)
    add_header(
        slide,
        PAGE_TITLES[6],
        "现有数据与推理 API 可封装为智能体工具；嵌入底座负责提供可复用地理表示",
    )
    _add_panel_frame(slide, 0.5, 1.35, 7.15, 4.85)
    add_text(slide, "自然语言需求到地图结果的协同流程", 0.73, 1.52, 3.85, 0.25, 16, DEEP_BLUE, True)
    add_text(slide, "概念流程示意（非真实系统界面）", 4.78, 1.53, 2.55, 0.24, 10, RED, True)
    add_picture(slide, WORKFLOW_IMAGE, 0.7, 1.88, 6.75, 4.02, "contain")

    tool_rows = [
        ("理解任务", "区域、月份、目标类别与输出形式"),
        ("读取嵌入", "embedding_tool：加载月度 64 维嵌入"),
        ("选择工具", "系统任务头、变化分析、任务摘要"),
        ("空间计算", "全域推理、统计汇总与质量检查"),
        ("交付结果", "地图图层、案例图与结构化报告"),
    ]
    for index, (title, detail) in enumerate(tool_rows, start=1):
        y = 1.35 + (index - 1) * 0.84
        add_rect(slide, 8.0, y, 4.78, 0.76, PALE_BLUE, LINE, MSO_SHAPE.ROUNDED_RECTANGLE)
        _add_number_badge(slide, index, 8.2, y + 0.19)
        add_text(slide, title, 8.75, y + 0.11, 3.5, 0.22, 14, DEEP_BLUE, True)
        add_text(slide, detail, 8.75, y + 0.41, 3.7, 0.18, 11, MUTED)

    add_rect(slide, 8.0, 5.64, 4.78, 0.58, PALE_GRAY, PALE_GRAY, MSO_SHAPE.ROUNDED_RECTANGLE)
    add_text(
        slide,
        "边界：智能体调用模型与工具，不替代模型推理；\n关键结果仍需地图复核。",
        8.2,
        5.76,
        4.4,
        0.34,
        10,
        RED,
        True,
    )
    add_text(
        slide,
        "当前状态：具备数据、推理与结构化分析接口；正式智能体编排属于下一步集成工作。",
        0.55,
        6.56,
        12.15,
        0.3,
        15,
        INK,
        True,
    )
    return slide


def build_page07(prs: Presentation):
    """Build the missing-observation page with independently embedded time panels."""
    slide = new_blank_slide(prs)
    add_header(
        slide,
        PAGE_TITLES[7],
        "融合目标时刻前后的多源观测，为缺测场景提供可解释的参考影像",
    )
    _set_header_title_size(slide, 24)
    add_text(
        slide,
        "时序输入：五个独立面板（源条裁去底部标签区）",
        0.52,
        1.05,
        5.2,
        0.22,
        13,
        MUTED,
        True,
    )
    panel_titles = ("前期 S2", "后期 S2", "变化前嵌入", "变化后嵌入", "变化概率")
    for index, title in enumerate(panel_titles):
        x = 0.5 + index * 2.48
        _add_panel_frame(slide, x, 1.35, 2.35, 2.61)
        add_text(slide, title, x + 0.1, 1.57, 2.15, 0.22, 13, DEEP_BLUE, True)
        left = index * 275
        right = left + 275 if index < 4 else 1376
        _add_cropped_picture(
            slide,
            OBSERVATION_TIMELINE,
            (left, 0, right, 262),
            x + 0.15,
            1.98,
            2.05,
            2.05 * 262 / (right - left),
        )
        if index < len(panel_titles) - 1:
            add_rect(slide, x + 2.32, 2.82, 0.2, 0.22, PALE_BLUE, PALE_BLUE, MSO_SHAPE.RIGHT_ARROW)

    add_rect(slide, 0.5, 4.25, 4.75, 1.72, PALE_BLUE, LINE, MSO_SHAPE.ROUNDED_RECTANGLE)
    add_text(slide, "时序面板说明", 0.76, 4.52, 2.1, 0.22, 16, DEEP_BLUE, True)
    add_text(
        slide,
        "前两幅展示相邻观测与目标时刻缺失；后续三幅为\n"
        "嵌入与变化候选线索，用于解释模型输入与输出边界。",
        0.76,
        4.88,
        4.05,
        0.5,
        12,
        MUTED,
    )
    add_rect(slide, 5.18, 4.87, 0.42, 0.3, PALE_BLUE, PALE_BLUE, MSO_SHAPE.RIGHT_ARROW)
    add_text(slide, "独立样例", 5.03, 4.48, 0.72, 0.2, 11, MUTED, True)

    output_x = 5.85
    _add_panel_frame(slide, output_x, 4.25, 6.93, 1.72)
    add_text(
        slide,
        "真实运行输出 | patch_000093",
        output_x + 2.52,
        4.52,
        3.9,
        0.22,
        16,
        DEEP_BLUE,
        True,
    )
    add_picture(slide, RECON_EXAMPLE, output_x + 0.2, 4.52, 2.08, 1.16, "cover")
    add_text(
        slide,
        "平台已保存生成结果 | 目标日期元数据未公开",
        output_x + 2.52,
        4.93,
        3.92,
        0.2,
        12,
        MUTED,
    )
    add_text(
        slide,
        "输出为参考影像，不等同于该时刻的真实观测。",
        output_x + 2.52,
        5.28,
        3.92,
        0.18,
        11,
        RED,
        True,
    )

    add_rect(slide, 0.5, 6.24, 12.28, 0.78, PALE_GRAY, PALE_GRAY, MSO_SHAPE.ROUNDED_RECTANGLE)
    add_text(
        slide,
        "使用边界｜左侧两图用于说明缺测场景，右图为独立真实运行样例，"
        "三者不构成同一 patch 的对照；\n"
        "模型输出是“参考影像”，不是该时刻真实观测，重要结论仍需真实影像核验。",
        0.72,
        6.38,
        11.85,
        0.46,
        11,
        RED,
        True,
    )
    return slide


def build_inspection_deck(output: Path) -> Path:
    """Save a four-slide inspection deck and confirm its package can be reopened."""
    presentation = Presentation()
    for builder in (build_page04, build_page05, build_page06, build_page07):
        builder(presentation)
    output.parent.mkdir(parents=True, exist_ok=True)
    presentation.save(output)
    reopened = Presentation(output)
    if len(reopened.slides) != 4:
        raise RuntimeError(f"Expected four inspection slides, got {len(reopened.slides)}")
    return output


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--output",
        type=Path,
        default=PRESENTATION_ROOT / "浦江交流_第04至07页_可编辑检查版_20260726.pptx",
    )
    return parser.parse_args()


def main() -> None:
    print(build_inspection_deck(parse_args().output))


if __name__ == "__main__":
    main()
