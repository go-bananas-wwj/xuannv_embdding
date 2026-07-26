#!/usr/bin/env python3
"""Build the first Pujiang presentation slide and its reproducible map assets."""

from __future__ import annotations

import argparse
from concurrent.futures import ThreadPoolExecutor
import io
import json
from dataclasses import dataclass
from pathlib import Path
import time
from typing import Callable
from urllib.parse import urlencode
from urllib.request import urlopen

import numpy as np
import rasterio
import torch
from PIL import Image, ImageDraw, ImageFont, ImageOps
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt
from sklearn.decomposition import PCA
from torch import nn

BLUE = RGBColor(0, 97, 170)
INK = RGBColor(25, 40, 55)
MUTED = RGBColor(81, 101, 118)
LINE = RGBColor(205, 215, 223)
PALE_BLUE = RGBColor(236, 246, 252)
WHITE = RGBColor(255, 255, 255)
FONT = "Microsoft YaHei"


@dataclass(frozen=True)
class PatchLayout:
    patch_id: str
    row: int
    col: int


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--source-pptx",
        type=Path,
        default=Path(
            "/root/.codex/attachments/d0a9024d-03de-463f-83fa-675e2a235467/" "Alpha Earth介绍.pptx"
        ),
    )
    parser.add_argument(
        "--api-base",
        default="http://60.31.21.42:22065",
    )
    parser.add_argument(
        "--api-cache-root",
        type=Path,
        default=Path(
            "/data/xuannv_embedding/experiments/presentation_pujiang_202607/api_assets"
        ),
    )
    parser.add_argument(
        "--asset-root",
        type=Path,
        default=Path("/root/workspace/xuannv/docs/presentations/assets/pujiang_202607/slide01"),
    )
    parser.add_argument(
        "--output-pptx",
        type=Path,
        default=Path(
            "/root/workspace/xuannv/docs/presentations/pujiang_202607/"
            "玄女月度地理嵌入_浦江交流_第01页_20260726.pptx"
        ),
    )
    parser.add_argument("--month", default="202605")
    parser.add_argument("--haidian-month", default="202605")
    parser.add_argument("--pca-sample-pixels", type=int, default=220_000)
    parser.add_argument("--seed", type=int, default=42)
    return parser.parse_args()


def api_url(base: str, path: str, **query: str | int) -> str:
    suffix = f"?{urlencode(query)}" if query else ""
    return f"{base.rstrip('/')}/{path.lstrip('/')}{suffix}"


def read_url_with_retry(
    url: str,
    attempts: int = 5,
    opener: Callable = urlopen,
    sleep: Callable[[float], None] = time.sleep,
) -> bytes:
    for attempt in range(attempts):
        try:
            with opener(url, timeout=120) as response:
                return response.read()
        except OSError:
            if attempt + 1 == attempts:
                raise
            sleep(2**attempt)
    raise RuntimeError("unreachable")


def api_get_json(base: str, path: str, **query: str | int) -> dict:
    payload = read_url_with_retry(api_url(base, path, **query))
    return json.loads(payload)


def api_download(base: str, path: str, output: Path, **query: str | int) -> str:
    url = api_url(base, path, **query)
    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_bytes(read_url_with_retry(url))
    return url


def api_patches(base: str, region: str) -> list[dict]:
    first = api_get_json(base, f"regions/{region}/patches", page=1, page_size=100)
    patches = list(first["patches"])
    page = 2
    while len(patches) < int(first["total"]):
        payload = api_get_json(
            base,
            f"regions/{region}/patches",
            page=page,
            page_size=100,
        )
        patches.extend(payload["patches"])
        page += 1
    if len(patches) != int(first["total"]):
        raise RuntimeError(f"{region} API patch list is incomplete")
    return patches


def api_patch_layout(patches: list[dict]) -> tuple[list[PatchLayout], int, int]:
    lefts = sorted({round(float(item["bounds"][0]), 3) for item in patches})
    tops = sorted(
        {round(float(item["bounds"][3]), 3) for item in patches},
        reverse=True,
    )
    col_of = {value: index for index, value in enumerate(lefts)}
    row_of = {value: index for index, value in enumerate(tops)}
    layouts = [
        PatchLayout(
            patch_id=str(item["patch_id"]),
            row=row_of[round(float(item["bounds"][3]), 3)],
            col=col_of[round(float(item["bounds"][0]), 3)],
        )
        for item in patches
    ]
    return layouts, len(tops), len(lefts)


def binary_prediction_mosaic(
    layouts: list[PatchLayout],
    rows: int,
    cols: int,
    predictions: dict[str, np.ndarray],
    threshold: float,
) -> tuple[np.ndarray, dict[str, float | int]]:
    first = np.asarray(predictions[layouts[0].patch_id]).squeeze()
    if first.ndim != 2:
        raise ValueError(f"Expected 2D prediction, got {first.shape}")
    tile_h, tile_w = first.shape
    canvas = np.full((rows * tile_h, cols * tile_w, 3), 255, dtype=np.uint8)
    positive = total = 0
    for layout in layouts:
        prediction = np.asarray(predictions[layout.patch_id]).squeeze()
        if prediction.shape != (tile_h, tile_w):
            raise ValueError(f"Inconsistent prediction shape for {layout.patch_id}")
        mask = np.nan_to_num(prediction, nan=0.0) >= threshold
        tile = np.full((tile_h, tile_w, 3), 255, dtype=np.uint8)
        tile[mask] = (230, 0, 0)
        y0, x0 = layout.row * tile_h, layout.col * tile_w
        canvas[y0 : y0 + tile_h, x0 : x0 + tile_w] = tile
        positive += int(mask.sum())
        total += int(mask.size)
    return canvas, {
        "patches": len(layouts),
        "threshold": threshold,
        "positive_pixel_ratio": positive / total,
    }


def build_api_binary_task_mosaic(
    base: str,
    region: str,
    task: str,
    version: str,
    month: str,
    cache_root: Path,
    output: Path,
) -> dict[str, float | int | str]:
    summary = api_get_json(
        base,
        f"regions/{region}/tasks/{task}/summary",
        version=version,
        month=month,
    )
    threshold = float(summary.get("prediction_statistics", {}).get("threshold") or 0.5)
    patches = api_patches(base, region)
    layouts, rows, cols = api_patch_layout(patches)
    prediction_root = cache_root / region / version / task / month
    prediction_root.mkdir(parents=True, exist_ok=True)

    def fetch(patch_id: str) -> tuple[str, np.ndarray]:
        target = prediction_root / f"{patch_id}.npy"
        url = api_url(
            base,
            f"regions/{region}/patches/{patch_id}/tasks/{task}/prediction",
            version=version,
            period=month,
        )
        if not target.exists():
            target.write_bytes(read_url_with_retry(url))
        return patch_id, np.load(target, allow_pickle=False)

    with ThreadPoolExecutor(max_workers=6) as pool:
        predictions = dict(pool.map(fetch, [item.patch_id for item in layouts]))
    canvas, stats = binary_prediction_mosaic(
        layouts,
        rows,
        cols,
        predictions,
        threshold,
    )
    save_slide_image(Image.fromarray(canvas), output)
    return {
        **stats,
        "region": region,
        "task": task,
        "version": version,
        "month": month,
        "head": summary["model"]["head_type"],
        "feature_source": summary["model"]["feature_source"],
        "api_summary": api_url(
            base,
            f"regions/{region}/tasks/{task}/summary",
            version=version,
            month=month,
        ),
        "api_prediction_template": api_url(
            base,
            f"regions/{region}/patches/{{patch_id}}/tasks/{task}/prediction",
            version=version,
            period=month,
        ),
    }


def font(size: int, bold: bool = False) -> ImageFont.FreeTypeFont:
    name = "NotoSansCJK-Bold.ttc" if bold else "NotoSansCJK-Regular.ttc"
    candidates = [
        Path("/usr/share/fonts/google-noto-cjk") / name,
        Path("/usr/share/fonts/opentype/noto") / name,
        Path("/usr/share/fonts/noto-cjk") / name,
    ]
    for path in candidates:
        if path.exists():
            return ImageFont.truetype(str(path), size)
    return ImageFont.load_default(size=size)


def load_layout(label_root: Path) -> tuple[list[PatchLayout], int, int, int, int]:
    paths = sorted((label_root / "masks").glob("patch_*.tif"))
    if not paths:
        raise FileNotFoundError(f"No masks found under {label_root / 'masks'}")
    records: list[tuple[str, float, float]] = []
    lefts: list[float] = []
    tops: list[float] = []
    tile_h = tile_w = 0
    for path in paths:
        with rasterio.open(path) as src:
            left, top = float(src.bounds.left), float(src.bounds.top)
            tile_h, tile_w = src.height, src.width
        records.append((path.stem, left, top))
        lefts.append(left)
        tops.append(top)
    unique_lefts = sorted({round(value, 3) for value in lefts})
    unique_tops = sorted({round(value, 3) for value in tops}, reverse=True)
    col_of = {value: index for index, value in enumerate(unique_lefts)}
    row_of = {value: index for index, value in enumerate(unique_tops)}
    layouts = [
        PatchLayout(
            patch_id=patch_id,
            row=row_of[round(top, 3)],
            col=col_of[round(left, 3)],
        )
        for patch_id, left, top in records
    ]
    return layouts, len(unique_tops), len(unique_lefts), tile_h, tile_w


def embedding_path(root: Path, patch_id: str, month: str) -> Path:
    return root / patch_id / f"{month}_embedding_map.pt"


def load_embedding(root: Path, patch_id: str, month: str) -> np.ndarray:
    path = embedding_path(root, patch_id, month)
    if not path.exists():
        raise FileNotFoundError(path)
    tensor = torch.load(path, map_location="cpu", weights_only=True)
    if tensor.shape != (64, 128, 128):
        raise ValueError(f"Unexpected embedding shape {tuple(tensor.shape)} at {path}")
    return tensor.float().numpy()


def stretch_rgb(values: np.ndarray, valid: np.ndarray) -> np.ndarray:
    result = np.zeros_like(values, dtype=np.float32)
    for channel in range(3):
        band = values[..., channel]
        lo, hi = np.percentile(band[valid], [2, 98])
        if hi > lo:
            result[..., channel] = np.clip((band - lo) / (hi - lo), 0.0, 1.0)
    result[~valid] = 0.94
    return (result * 255).astype(np.uint8)


def save_slide_image(image: Image.Image, output: Path, max_side: int = 1800) -> None:
    image = image.convert("RGB")
    image.thumbnail((max_side, max_side), Image.Resampling.LANCZOS)
    output.parent.mkdir(parents=True, exist_ok=True)
    image.save(output, optimize=True)


def build_harbin_pca(
    root: Path,
    label_root: Path,
    month: str,
    output: Path,
    max_pixels: int,
    seed: int,
) -> dict[str, int | str]:
    layouts, rows, cols, tile_h, tile_w = load_layout(label_root)
    missing = [
        layout.patch_id
        for layout in layouts
        if not embedding_path(root, layout.patch_id, month).exists()
    ]
    if missing:
        raise RuntimeError(f"Harbin embedding export is incomplete: {len(missing)} patches missing")

    rng = np.random.default_rng(seed)
    per_patch = max(64, max_pixels // len(layouts))
    samples: list[np.ndarray] = []
    for layout in layouts:
        emb = load_embedding(root, layout.patch_id, month)
        flat = emb.reshape(64, -1).T
        indices = rng.choice(flat.shape[0], size=min(per_patch, flat.shape[0]), replace=False)
        samples.append(flat[indices])
    pca = PCA(n_components=3, random_state=seed).fit(np.concatenate(samples, axis=0))

    canvas = np.zeros((rows * tile_h, cols * tile_w, 3), dtype=np.float32)
    valid = np.zeros(canvas.shape[:2], dtype=bool)
    for layout in layouts:
        emb = load_embedding(root, layout.patch_id, month)
        rgb = pca.transform(emb.reshape(64, -1).T).reshape(tile_h, tile_w, 3)
        y0, x0 = layout.row * tile_h, layout.col * tile_w
        canvas[y0 : y0 + tile_h, x0 : x0 + tile_w] = rgb
        valid[y0 : y0 + tile_h, x0 : x0 + tile_w] = True

    save_slide_image(Image.fromarray(stretch_rgb(canvas, valid)), output)
    return {
        "patches": len(layouts),
        "rows": rows,
        "cols": cols,
        "tile_height": tile_h,
        "tile_width": tile_w,
        "month": month,
    }


def clean_existing_asset(source: Path, output: Path, header_pixels: int = 70) -> None:
    with Image.open(source) as image:
        image = image.convert("RGB")
        if header_pixels:
            image = image.crop((0, header_pixels, image.width, image.height))
        save_slide_image(image, output)


def clean_harbin_task_map(source: Path, output: Path) -> None:
    with Image.open(source) as image:
        image = image.convert("RGB")
        # Keep the plotted geographic extent and remove the English title and axes.
        image = image.crop((125, 55, 1178, 1112))
        image = ImageOps.expand(image, border=4, fill=(215, 222, 228))
        save_slide_image(image, output)


class LegacyStage2UNetHead(nn.Module):
    """Two-stage lightweight UNet head used by the original Harbin evaluation."""

    def __init__(self, embed_dim: int = 64, num_classes: int = 2) -> None:
        super().__init__()
        self.up1 = nn.ConvTranspose2d(embed_dim, embed_dim // 2, kernel_size=2, stride=2)
        self.conv1 = nn.Sequential(
            nn.Conv2d(
                embed_dim + embed_dim // 2,
                embed_dim // 2,
                kernel_size=3,
                padding=1,
            ),
            nn.BatchNorm2d(embed_dim // 2),
            nn.ReLU(inplace=True),
        )
        self.up2 = nn.ConvTranspose2d(embed_dim // 2, embed_dim // 4, kernel_size=2, stride=2)
        self.conv2 = nn.Sequential(
            nn.Conv2d(
                embed_dim + embed_dim // 4,
                embed_dim // 4,
                kernel_size=3,
                padding=1,
            ),
            nn.BatchNorm2d(embed_dim // 4),
            nn.ReLU(inplace=True),
        )
        self.final = nn.Conv2d(embed_dim // 4, num_classes, kernel_size=1)

    def forward(self, value: torch.Tensor) -> torch.Tensor:
        first = self.up1(value)
        first_skip = nn.functional.interpolate(
            value, scale_factor=2, mode="bilinear", align_corners=False
        )
        first = self.conv1(torch.cat([first, first_skip], dim=1))
        second = self.up2(first)
        second_skip = nn.functional.interpolate(
            value, scale_factor=4, mode="bilinear", align_corners=False
        )
        second = self.conv2(torch.cat([second, second_skip], dim=1))
        logits = self.final(second)
        return nn.functional.interpolate(
            logits, size=value.shape[-2:], mode="bilinear", align_corners=False
        )


def build_harbin_prediction(
    embedding_root: Path,
    label_root: Path,
    month: str,
    checkpoint: Path,
    metrics_path: Path,
    output: Path,
) -> dict[str, float | int | str | None]:
    state = torch.load(checkpoint, map_location="cpu", weights_only=True)
    model = LegacyStage2UNetHead()
    model.load_state_dict(state)
    model.eval()
    metrics = json.loads(metrics_path.read_text(encoding="utf-8"))
    threshold = float(metrics["best_threshold"])
    layouts, rows, cols, tile_h, tile_w = load_layout(label_root)
    canvas = np.full((rows * tile_h, cols * tile_w, 3), 240, dtype=np.uint8)
    batch_size = 4
    with torch.no_grad():
        for start in range(0, len(layouts), batch_size):
            batch_layouts = layouts[start : start + batch_size]
            batch = torch.stack(
                [
                    torch.from_numpy(load_embedding(embedding_root, layout.patch_id, month))
                    for layout in batch_layouts
                ]
            )
            probabilities = torch.sigmoid(model(batch)[:, 1]).numpy()
            for layout, probability in zip(batch_layouts, probabilities, strict=True):
                prediction = probability >= threshold
                tile = np.full((tile_h, tile_w, 3), 255, dtype=np.uint8)
                tile[prediction] = (226, 38, 49)
                y0, x0 = layout.row * tile_h, layout.col * tile_w
                canvas[y0 : y0 + tile_h, x0 : x0 + tile_w] = tile
    save_slide_image(Image.fromarray(canvas), output)
    return {
        "checkpoint": str(checkpoint),
        "metrics": str(metrics_path),
        "threshold": threshold,
        "head": "legacy_stage2_lightweight_unet",
        "fold": int(metrics["fold"]),
        "best_epoch": int(metrics["best_epoch"]),
        "f1": float(metrics["f1_best"]),
        "auc": float(metrics["auc_roc"]) if "auc_roc" in metrics else None,
    }


def clear_slide(slide) -> None:
    for shape in list(slide.shapes):
        shape._element.getparent().remove(shape._element)


def keep_only_slide(prs: Presentation, index: int) -> None:
    slide_ids = prs.slides._sldIdLst
    keep = slide_ids[index]
    for slide_id in list(slide_ids):
        if slide_id is not keep:
            slide_ids.remove(slide_id)


def add_rect(slide, x: float, y: float, w: float, h: float, fill, line=None):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = line or fill
    return shape


def add_text(
    slide,
    value: str,
    x: float,
    y: float,
    w: float,
    h: float,
    size: int,
    color=INK,
    bold: bool = False,
    align=PP_ALIGN.LEFT,
):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    paragraph.space_after = Pt(0)
    run = paragraph.add_run()
    run.text = value
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return shape


def add_picture_contain(slide, path: Path, x: float, y: float, w: float, h: float):
    with Image.open(path) as image:
        ratio = image.width / image.height
    box_ratio = w / h
    if ratio >= box_ratio:
        draw_w, draw_h = w, w / ratio
    else:
        draw_h, draw_w = h, h * ratio
    draw_x = x + (w - draw_w) / 2
    draw_y = y + (h - draw_h) / 2
    return slide.shapes.add_picture(
        str(path), Inches(draw_x), Inches(draw_y), Inches(draw_w), Inches(draw_h)
    )


def add_region(
    slide,
    name: str,
    x: float,
    pca_path: Path,
    tasks: list[tuple[str, Path]],
    footer: str,
    pca_caption: str,
) -> None:
    add_text(slide, name, x, 1.60, 6.05, 0.34, 18, BLUE, True, PP_ALIGN.CENTER)
    add_text(
        slide,
        pca_caption,
        x + 0.05,
        1.94,
        3.82,
        0.25,
        14,
        INK,
        True,
    )
    add_rect(slide, x + 0.04, 2.20, 3.86, 3.93, WHITE, LINE)
    add_picture_contain(slide, pca_path, x + 0.10, 2.26, 3.74, 3.81)
    single_task = len(tasks) == 1
    for index, (label, path) in enumerate(tasks):
        y = 2.20 + index * 1.31
        image_height = 3.65 if single_task else 0.94
        add_text(
            slide,
            label,
            x + 4.08,
            y,
            1.78,
            0.25,
            14,
            INK,
            True,
            PP_ALIGN.CENTER,
        )
        add_rect(slide, x + 4.08, y + 0.28, 1.78, image_height, WHITE, LINE)
        add_picture_contain(
            slide,
            path,
            x + 4.13,
            y + 0.33,
            1.68,
            image_height - 0.10,
        )
    add_text(
        slide,
        "■ 红=预测目标｜白=背景",
        x + 4.02,
        6.10,
        1.92,
        0.32,
        9,
        RGBColor(180, 35, 45),
        True,
        PP_ALIGN.CENTER,
    )
    add_rect(slide, x + 0.04, 6.42, 5.82, 0.53, PALE_BLUE, PALE_BLUE)
    add_text(
        slide,
        footer,
        x + 0.12,
        6.44,
        5.66,
        0.46,
        10,
        MUTED,
        False,
        PP_ALIGN.CENTER,
    )


def build_ppt(source: Path, output: Path, assets: dict[str, Path]) -> None:
    prs = Presentation(source)
    keep_only_slide(prs, 1)
    slide = prs.slides[0]
    clear_slide(slide)
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = WHITE

    add_rect(slide, 0.0, 0.0, 3.18, 0.48, BLUE)
    add_text(slide, "玄女月度地理嵌入", 0.35, 0.02, 2.55, 0.42, 17, WHITE, True)
    add_rect(slide, 0.0, 0.48, 13.333, 0.68, BLUE)
    add_text(
        slide,
        "玄女月度地理嵌入：海淀区与哈尔滨新区实践",
        0.42,
        0.50,
        12.5,
        0.62,
        29,
        WHITE,
        True,
        PP_ALIGN.CENTER,
    )
    add_text(
        slide,
        "两个区域已形成月度、稠密、可复用的地理嵌入，同一份嵌入支持多类地物制图任务",
        0.45,
        1.18,
        12.43,
        0.34,
        20,
        INK,
        True,
        PP_ALIGN.CENTER,
    )
    add_rect(slide, 6.64, 1.62, 0.015, 5.29, LINE, LINE)
    add_region(
        slide,
        "海淀区｜P10C epoch 800",
        0.31,
        assets["haidian_pca"],
        [
            ("建筑提取", assets["haidian_building"]),
            ("道路提取", assets["haidian_road"]),
            ("水体提取", assets["haidian_water"]),
        ],
        "320 patch｜月度：2025.12–2026.05｜128×128×64 / patch",
        "2026年5月全域 Embedding PCA",
    )
    add_region(
        slide,
        "哈尔滨新区｜V5（API v2）",
        7.16,
        assets["harbin_pca"],
        [
            ("建设区域监测", assets["harbin_building"]),
        ],
        "424 patch｜10期：2025.04–2026.05｜64维月度嵌入",
        "2026年5月全域 Embedding PCA",
    )
    add_text(
        slide,
        "注：PCA 颜色表示 64 维嵌入的前三主成分；不同区域的颜色本身不可直接对应。",
        0.48,
        7.08,
        12.35,
        0.22,
        14,
        MUTED,
        False,
        PP_ALIGN.CENTER,
    )
    output.parent.mkdir(parents=True, exist_ok=True)
    prs.save(output)


def contain_pil(canvas: Image.Image, source: Path, box: tuple[int, int, int, int]) -> None:
    x0, y0, x1, y1 = box
    with Image.open(source) as image:
        image = image.convert("RGB")
        image.thumbnail((x1 - x0, y1 - y0), Image.Resampling.LANCZOS)
        x = x0 + (x1 - x0 - image.width) // 2
        y = y0 + (y1 - y0 - image.height) // 2
        canvas.paste(image, (x, y))


def build_preview(output: Path, assets: dict[str, Path]) -> None:
    canvas = Image.new("RGB", (1600, 900), "white")
    draw = ImageDraw.Draw(canvas)
    blue = (0, 97, 170)
    ink = (25, 40, 55)
    muted = (81, 101, 118)
    line = (205, 215, 223)
    pale = (236, 246, 252)

    draw.rectangle((0, 0, 382, 58), fill=blue)
    draw.text((42, 11), "玄女月度地理嵌入", font=font(25, True), fill="white")
    draw.rectangle((0, 58, 1600, 139), fill=blue)
    title = "玄女月度地理嵌入：海淀区与哈尔滨新区实践"
    title_box = draw.textbbox((0, 0), title, font=font(36, True))
    draw.text(
        ((1600 - (title_box[2] - title_box[0])) // 2, 77), title, font=font(36, True), fill="white"
    )
    conclusion = "两个区域已形成月度、稠密、可复用的地理嵌入，同一份嵌入支持多类地物制图任务"
    box = draw.textbbox((0, 0), conclusion, font=font(19, True))
    draw.text(((1600 - (box[2] - box[0])) // 2, 147), conclusion, font=font(19, True), fill=ink)
    draw.line((798, 194, 798, 829), fill=line, width=2)

    regions = [
        (
            37,
            "海淀区｜P10C epoch 800",
            assets["haidian_pca"],
            [
                ("建筑提取", assets["haidian_building"]),
                ("道路提取", assets["haidian_road"]),
                ("水体提取", assets["haidian_water"]),
            ],
            "320 patch｜月度：2025.12–2026.05｜128×128×64 / patch",
            "2026年5月全域 Embedding PCA",
        ),
        (
            857,
            "哈尔滨新区｜V5（API v2）",
            assets["harbin_pca"],
            [
                ("建设区域监测", assets["harbin_building"]),
            ],
            "424 patch｜10期：2025.04–2026.05｜64维月度嵌入",
            "2026年5月全域 Embedding PCA",
        ),
    ]
    for x, name, pca_path, tasks, footer, pca_caption in regions:
        name_box = draw.textbbox((0, 0), name, font=font(23, True))
        draw.text(
            (x + (726 - (name_box[2] - name_box[0])) // 2, 189),
            name,
            font=font(23, True),
            fill=blue,
        )
        draw.text((x + 7, 232), pca_caption, font=font(16, True), fill=ink)
        draw.rectangle((x + 5, 264, x + 469, 736), outline=line, width=2)
        contain_pil(canvas, pca_path, (x + 12, 271, x + 462, 729))
        single_task = len(tasks) == 1
        for index, (label, path) in enumerate(tasks):
            y = 264 + index * 157
            image_bottom = y + 438 if single_task else y + 146
            label_box = draw.textbbox((0, 0), label, font=font(15, True))
            draw.text(
                (x + 492 + (214 - (label_box[2] - label_box[0])) // 2, y),
                label,
                font=font(15, True),
                fill=ink,
            )
            draw.rectangle((x + 491, y + 34, x + 706, image_bottom), outline=line, width=2)
            contain_pil(canvas, path, (x + 497, y + 40, x + 700, image_bottom - 6))
        legend = "■ 红=预测目标｜白=背景"
        legend_box = draw.textbbox((0, 0), legend, font=font(11, True))
        draw.text(
            (x + 491 + (215 - (legend_box[2] - legend_box[0])) // 2, 731),
            legend,
            font=font(11, True),
            fill=(180, 35, 45),
        )
        draw.rectangle((x + 5, 770, x + 706, 834), fill=pale)
        footer_box = draw.textbbox((0, 0), footer, font=font(12))
        draw.text(
            (x + (711 - (footer_box[2] - footer_box[0])) // 2, 791),
            footer,
            font=font(12),
            fill=muted,
        )

    note = "注：PCA 颜色表示 64 维嵌入的前三主成分；不同区域的颜色本身不可直接对应。"
    note_box = draw.textbbox((0, 0), note, font=font(14))
    draw.text(
        ((1600 - (note_box[2] - note_box[0])) // 2, 859),
        note,
        font=font(14),
        fill=muted,
    )
    output.parent.mkdir(parents=True, exist_ok=True)
    canvas.save(output)


def main() -> None:
    args = parse_args()
    args.asset_root.mkdir(parents=True, exist_ok=True)
    sources = {
        "haidian_pca": Path(
            "/data/xuannv_embedding/experiments/presentation_pujiang_202607/"
            "haidian_p10c_epoch800_202605_full_domain/P10C_embedding_pca_202605_geo.png"
        ),
        "haidian_building": Path(
            "/data/xuannv_embedding/experiments/presentation_pujiang_202607/"
            "haidian_p10c_epoch800_202605_full_domain/building/"
            "P10C_building_prediction_geo.png"
        ),
        "haidian_road": Path(
            "/data/xuannv_embedding/experiments/presentation_pujiang_202607/"
            "haidian_p10c_epoch800_202605_full_domain/road/P10C_road_prediction_geo.png"
        ),
        "haidian_water": Path(
            "/data/xuannv_embedding/experiments/presentation_pujiang_202607/"
            "haidian_p10c_epoch800_202605_full_domain/water/P10C_water_prediction_geo.png"
        ),
    }
    assets = {
        key: args.asset_root / f"{key}.png"
        for key in [
            *sources,
            "harbin_pca",
            "harbin_building",
        ]
    }
    for key, source in sources.items():
        if key in {"haidian_building", "haidian_road", "haidian_water"}:
            continue
        clean_existing_asset(source, assets[key])

    harbin_pca_cache = (
        args.api_cache_root / f"harbin_v5_embedding_{args.month}_api.png"
    )
    harbin_pca_url = api_download(
        args.api_base,
        "regions/harbin/mosaic",
        harbin_pca_cache,
        date=args.month,
        sensor_type="embedding",
        version="v2",
        format="png",
    )
    clean_existing_asset(harbin_pca_cache, assets["harbin_pca"], header_pixels=0)
    haidian_building_meta = build_api_binary_task_mosaic(
        args.api_base,
        "haidian",
        "building_extraction",
        "v1",
        args.haidian_month,
        args.api_cache_root,
        assets["haidian_building"],
    )
    haidian_road_meta = build_api_binary_task_mosaic(
        args.api_base,
        "haidian",
        "road_extraction",
        "v1",
        args.haidian_month,
        args.api_cache_root,
        assets["haidian_road"],
    )
    haidian_water_meta = build_api_binary_task_mosaic(
        args.api_base,
        "haidian",
        "water_extraction",
        "v1",
        args.haidian_month,
        args.api_cache_root,
        assets["haidian_water"],
    )
    harbin_building_meta = build_api_binary_task_mosaic(
        args.api_base,
        "harbin",
        "building_extraction",
        "v2",
        args.month,
        args.api_cache_root,
        assets["harbin_building"],
    )
    pca_meta = {
        "patches": 424,
        "month": args.month,
        "version": "v2",
        "model": "V5",
        "api_url": harbin_pca_url,
        "cache_path": str(harbin_pca_cache),
    }
    probe_meta = {
        "haidian_building": haidian_building_meta,
        "haidian_road": haidian_road_meta,
        "haidian_water": haidian_water_meta,
        "harbin_building": harbin_building_meta,
    }
    build_ppt(args.source_pptx, args.output_pptx, assets)
    preview = args.asset_root / "slide01_preview.png"
    build_preview(preview, assets)
    metadata = {
        "slide": 1,
        "title": "玄女月度地理嵌入：海淀区与哈尔滨新区实践",
        "pptx": str(args.output_pptx),
        "preview": str(preview),
        "api_base": args.api_base,
        "haidian_model": {
            "experiment": "P10C",
            "checkpoint": "epoch_800",
            "embedding_version": "v1",
            "month": args.haidian_month,
            "building_head": "binary_conv3x3",
            "building_threshold": haidian_building_meta["threshold"],
            "building_task_map_type": "API numeric prediction thresholded with registered threshold",
            "building_source": haidian_building_meta["api_prediction_template"],
            "road_head": haidian_road_meta["head"],
            "road_threshold": haidian_road_meta["threshold"],
            "road_source": haidian_road_meta["api_prediction_template"],
            "water_head": haidian_water_meta["head"],
            "water_threshold": haidian_water_meta["threshold"],
            "water_source": haidian_water_meta["api_prediction_template"],
        },
        "harbin_model": {
            "experiment": "V5",
            "embedding_version": "v2",
            "month": args.month,
            "training": "deployed embedding-api asset",
            "embedding_source": harbin_pca_url,
            "task_map_type": "API construction-mapped numeric prediction",
            "task_label": "建设区域监测",
        },
        "harbin_pca": pca_meta,
        "api_task_maps": probe_meta,
        "assets": {key: str(path) for key, path in assets.items()},
    }
    (args.asset_root / "metadata.json").write_text(
        json.dumps(metadata, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )
    print(json.dumps(metadata, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
