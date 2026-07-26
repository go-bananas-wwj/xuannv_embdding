#!/usr/bin/env python3
"""Assemble the ten-slide editable Pujiang presentation."""

from __future__ import annotations

import argparse
import sys
from pathlib import Path

from pptx import Presentation

REPORT_DIR = Path(__file__).resolve().parent
if str(REPORT_DIR) not in sys.path:
    sys.path.insert(0, str(REPORT_DIR))

from pujiang_editable_common import clone_first_slide, is_full_slide_picture  # noqa: E402
from pujiang_editable_pages04_07 import (  # noqa: E402
    build_page04,
    build_page05,
    build_page06,
    build_page07,
)
from pujiang_editable_pages08_10 import build_page08, build_page09, build_page10  # noqa: E402

ROOT = Path(__file__).resolve().parents[2]
PRESENTATION_ROOT = ROOT / "docs/presentations/pujiang_202607"
OUTPUT_PATH = PRESENTATION_ROOT / "玄女月度地理嵌入_浦江交流_10页可编辑版_20260726.pptx"
SOURCE_SLIDES = tuple(
    PRESENTATION_ROOT / f"玄女月度地理嵌入_浦江交流_第{page:02d}页_20260726.pptx"
    for page in range(1, 4)
)


def build_deck(output: Path = OUTPUT_PATH) -> Path:
    """Clone editable slides 1-3, append native slides 4-10, save and reopen."""
    presentation = Presentation()
    for source in SOURCE_SLIDES:
        clone_first_slide(source, presentation)
    for builder in (
        build_page04,
        build_page05,
        build_page06,
        build_page07,
        build_page08,
        build_page09,
        build_page10,
    ):
        builder(presentation)
    if len(presentation.slides) != 10:
        raise RuntimeError(f"Expected ten slides before save, got {len(presentation.slides)}")
    if any(
        is_full_slide_picture(shape, presentation)
        for slide in presentation.slides
        for shape in slide.shapes
    ):
        raise RuntimeError("Editable deck contains a full-slide picture")

    output.parent.mkdir(parents=True, exist_ok=True)
    presentation.save(output)
    reopened = Presentation(output)
    if len(reopened.slides) != 10:
        raise RuntimeError(f"Expected ten slides after reopen, got {len(reopened.slides)}")
    return output


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", type=Path, default=OUTPUT_PATH)
    return parser.parse_args()


def main() -> None:
    print(build_deck(parse_args().output))


if __name__ == "__main__":
    main()
