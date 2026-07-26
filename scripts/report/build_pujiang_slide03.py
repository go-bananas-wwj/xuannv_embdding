#!/usr/bin/env python3
"""Build Pujiang presentation slide 3: AEF-inspired monthly city adaptations."""

from __future__ import annotations

import argparse
import json
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

DEFAULT_SOURCE_PPTX = Path(
    "/root/.codex/attachments/d0a9024d-03de-463f-83fa-675e2a235467/" "Alpha Earth介绍.pptx"
)
DEFAULT_ASSET_ROOT = Path("/root/workspace/xuannv/docs/presentations/assets/pujiang_202607/slide03")
DEFAULT_OUTPUT = Path(
    "/root/workspace/xuannv/docs/presentations/pujiang_202607/"
    "玄女月度地理嵌入_浦江交流_第03页_20260726.pptx"
)
DEFAULT_PREVIEW = DEFAULT_ASSET_ROOT / "slide03_preview.png"

BLUE = RGBColor(0, 97, 170)
DEEP_BLUE = RGBColor(0, 69, 122)
CYAN = RGBColor(0, 143, 174)
GREEN = RGBColor(34, 128, 94)
INK = RGBColor(25, 40, 55)
MUTED = RGBColor(81, 101, 118)
LINE = RGBColor(205, 215, 223)
PALE_BLUE = RGBColor(236, 246, 252)
PALE_CYAN = RGBColor(235, 248, 250)
PALE_GREEN = RGBColor(237, 247, 242)
PALE_GRAY = RGBColor(245, 247, 249)
WHITE = RGBColor(255, 255, 255)
RED = RGBColor(181, 45, 55)
FONT = "Microsoft YaHei"

BASE_NODES = [
    ("时序多源观测", "S2 / S1 / Landsat"),
    ("STP 时空编码", "空间 × 时间 × 精度"),
    ("64维嵌入场", "月度 128×128×64"),
    ("多目标重建", "跨源压缩学习"),
]
MONTHS = ["12月", "1月", "2月", "3月", "4月", "5月"]
MASK_FACTS = [
    ("月份 65%", "每个训练样本有 65% 概率启动月份遮挡\n随机拿走 1–4 个月"),
    ("空间 65%", "每个训练样本有 65% 概率启动空间遮挡\n遮住约 32% 的图像块"),
    ("补回完整答案", "答案仍是完整影像\n从剩余月份、其他传感器和周边空间恢复"),
    ("真实缺失不改", "availability mask 只记录原始数据是否真实存在"),
]


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source-pptx", type=Path, default=DEFAULT_SOURCE_PPTX)
    parser.add_argument("--asset-root", type=Path, default=DEFAULT_ASSET_ROOT)
    parser.add_argument("--output-pptx", type=Path, default=DEFAULT_OUTPUT)
    parser.add_argument("--preview", type=Path, default=DEFAULT_PREVIEW)
    return parser.parse_args()


def pil_font(size: int, bold: bool = False) -> ImageFont.FreeTypeFont:
    path = (
        "/usr/share/fonts/google-noto-cjk/NotoSansCJK-Bold.ttc"
        if bold
        else "/usr/share/fonts/google-noto-cjk/NotoSansCJK-Regular.ttc"
    )
    if Path(path).exists():
        return ImageFont.truetype(path, size)
    return ImageFont.truetype("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf", size)


def clear_slide(slide) -> None:
    for shape in list(slide.shapes):
        shape._element.getparent().remove(shape._element)


def keep_only_slide(prs: Presentation, index: int) -> None:
    slide_ids = prs.slides._sldIdLst
    keep = slide_ids[index]
    for slide_id in list(slide_ids):
        if slide_id is not keep:
            slide_ids.remove(slide_id)


def add_rect(slide, x: float, y: float, w: float, h: float, fill, line=None):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(x),
        Inches(y),
        Inches(w),
        Inches(h),
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = line or fill
    return shape


def add_text(
    slide,
    value: str,
    x: float,
    y: float,
    w: float,
    h: float,
    size: int,
    color=INK,
    bold: bool = False,
    align=PP_ALIGN.LEFT,
):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    paragraph.space_after = Pt(0)
    run = paragraph.add_run()
    run.text = value
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return shape


def add_header(slide) -> None:
    add_rect(slide, 0.0, 0.0, 3.18, 0.48, BLUE)
    add_text(slide, "玄女月度地理嵌入", 0.35, 0.02, 2.55, 0.42, 17, WHITE, True)
    add_rect(slide, 0.0, 0.48, 13.333, 0.68, BLUE)
    add_text(
        slide,
        "从 AEF 思路到城市尺度月度嵌入",
        0.42,
        0.50,
        12.5,
        0.62,
        29,
        WHITE,
        True,
        PP_ALIGN.CENTER,
    )
    add_text(
        slide,
        "参考多源压缩重建与 64 维嵌入场范式，针对月度云雾、小区域训练和细粒度目标定向适配",
        0.48,
        1.17,
        12.37,
        0.32,
        18,
        INK,
        True,
        PP_ALIGN.CENTER,
    )


def add_base_pipeline(slide) -> None:
    add_rect(slide, 0.48, 1.58, 12.37, 1.05, PALE_BLUE, LINE)
    add_text(
        slide,
        "共同底座｜参考 AEF 的 embedding field 范式",
        0.66,
        1.62,
        3.35,
        0.28,
        14,
        BLUE,
        True,
    )
    x0, y, width, gap = 0.70, 1.96, 2.65, 0.38
    for index, (title, detail) in enumerate(BASE_NODES):
        x = x0 + index * (width + gap)
        add_rect(slide, x, y, width, 0.50, WHITE, LINE)
        add_text(
            slide,
            title,
            x + 0.06,
            y + 0.02,
            width - 0.12,
            0.23,
            13,
            DEEP_BLUE,
            True,
            PP_ALIGN.CENTER,
        )
        add_text(
            slide,
            detail,
            x + 0.05,
            y + 0.24,
            width - 0.10,
            0.22,
            12,
            MUTED,
            False,
            PP_ALIGN.CENTER,
        )
        if index < len(BASE_NODES) - 1:
            add_text(slide, "→", x + width, y + 0.08, gap, 0.34, 16, BLUE, True, PP_ALIGN.CENTER)


def add_card_frame(slide, x: float, title: str, number: str, fill, accent) -> None:
    add_rect(slide, x, 2.83, 3.94, 3.78, fill, LINE)
    add_rect(slide, x, 2.83, 0.55, 0.57, accent, accent)
    add_text(slide, number, x, 2.86, 0.55, 0.50, 14, WHITE, True, PP_ALIGN.CENTER)
    add_text(slide, title, x + 0.69, 2.87, 3.05, 0.47, 18, accent, True)


def add_monthly_card(slide) -> None:
    x = 0.48
    add_card_frame(slide, x, "月度多源嵌入", "01", PALE_BLUE, BLUE)
    add_text(
        slide,
        "按月组织多源观测，不依赖年度合成",
        x + 0.22,
        3.47,
        3.50,
        0.35,
        14,
        INK,
        True,
        PP_ALIGN.CENTER,
    )
    line_y = 4.20
    add_rect(slide, x + 0.47, line_y + 0.14, 2.98, 0.025, BLUE, BLUE)
    for index, month in enumerate(MONTHS):
        cx = x + 0.50 + index * 0.58
        add_rect(slide, cx, line_y, 0.29, 0.29, BLUE, BLUE)
        add_text(
            slide,
            month,
            cx - 0.10,
            line_y + 0.34,
            0.49,
            0.25,
            12,
            MUTED,
            False,
            PP_ALIGN.CENTER,
        )
    add_rect(slide, x + 0.42, 5.05, 3.10, 0.57, WHITE, LINE)
    add_text(slide, "每期输出", x + 0.58, 5.09, 0.82, 0.47, 12, MUTED, True)
    add_text(slide, "128×128×64", x + 1.42, 5.09, 1.92, 0.47, 20, BLUE, True, PP_ALIGN.CENTER)
    add_text(
        slide,
        "连续 6 期月度索引｜跨月份注意力融合上下文",
        x + 0.27,
        5.78,
        3.40,
        0.53,
        12,
        INK,
        True,
        PP_ALIGN.CENTER,
    )


def add_masking_card(slide) -> None:
    x = 4.70
    add_card_frame(slide, x, "困难重建与缺失鲁棒", "02", PALE_CYAN, CYAN)
    add_text(
        slide,
        "先故意拿走一部分输入，再要求模型补回来",
        x + 0.22,
        3.43,
        3.50,
        0.34,
        13,
        INK,
        True,
        PP_ALIGN.CENTER,
    )
    add_text(
        slide,
        "整源模态遮挡｜S2 18%｜S1 / Landsat 35%",
        x + 0.24,
        3.79,
        3.46,
        0.30,
        13,
        CYAN,
        True,
        PP_ALIGN.CENTER,
    )
    y = 4.12
    for index, (title, detail) in enumerate(MASK_FACTS):
        row_y = y + index * 0.53
        add_rect(slide, x + 0.24, row_y, 0.10, 0.46, CYAN, CYAN)
        add_text(slide, title, x + 0.45, row_y, 1.06, 0.46, 14, CYAN, True)
        add_text(slide, detail, x + 1.50, row_y, 2.20, 0.46, 12, INK, True)
    add_rect(slide, x + 0.30, 6.25, 3.34, 0.26, WHITE, LINE)
    add_text(
        slide,
        "目的：学会在云雾、缺月或缺源时仍生成稳定嵌入",
        x + 0.38,
        6.27,
        3.18,
        0.22,
        13,
        DEEP_BLUE,
        True,
        PP_ALIGN.CENTER,
    )


def add_semantic_card(slide) -> None:
    x = 8.92
    add_card_frame(slide, x, "语义稳定与高分细节", "03", PALE_GREEN, GREEN)
    add_text(
        slide,
        "OSM 只做语义提示，不把下游真值写进底座",
        x + 0.25,
        3.47,
        3.44,
        0.35,
        14,
        INK,
        True,
        PP_ALIGN.CENTER,
    )
    add_text(
        slide,
        "13 项可重叠的 OSM 语义提示\n"
        "建筑、主路、支路、铁路、水体、绿地、农业、居住、商业、工业、施工、步道、操场",
        x + 0.25,
        3.84,
        3.44,
        0.70,
        12,
        INK,
        True,
        PP_ALIGN.CENTER,
    )
    add_rect(slide, x + 0.27, 4.56, 3.40, 0.76, WHITE, LINE)
    add_text(
        slide,
        "怎么进入模型？\n临时 1×1 线性探针对每个像素分别判断是否属于各项可重叠语义；"
        "训练结束后丢弃",
        x + 0.38,
        4.65,
        3.18,
        0.69,
        13,
        DEEP_BLUE,
        True,
        PP_ALIGN.CENTER,
    )
    add_text(
        slide,
        "0.14 控制整组 OSM 提示对训练的影响力度\n"
        "不是准确率、数据占比，也不代表总损失的 14%",
        x + 0.27,
        5.46,
        3.40,
        0.52,
        13,
        GREEN,
        True,
        PP_ALIGN.CENTER,
    )
    add_text(
        slide,
        "前 80 轮逐步加到 0.14",
        x + 0.27,
        6.03,
        3.40,
        0.34,
        13,
        MUTED,
        True,
        PP_ALIGN.CENTER,
    )


def build_ppt(source: Path, output: Path) -> None:
    prs = Presentation(source)
    keep_only_slide(prs, 4)
    slide = prs.slides[0]
    clear_slide(slide)
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = WHITE

    add_header(slide)
    add_base_pipeline(slide)
    add_monthly_card(slide)
    add_masking_card(slide)
    add_semantic_card(slide)
    add_rect(slide, 0.48, 6.76, 12.37, 0.37, PALE_GRAY, PALE_GRAY)
    add_text(
        slide,
        "证据边界｜不是完全复现 AEF；OSM 是辅助弱语义，不是人工真值；"
        "P10C 中人工变化监督 0、时间对比损失 0。",
        0.65,
        6.78,
        12.03,
        0.32,
        12,
        RED,
        True,
        PP_ALIGN.CENTER,
    )
    add_text(
        slide,
        "本页参数均来自海淀 P10C epoch 800 配置；不宣称全国或跨区域零样本泛化。",
        0.50,
        7.16,
        12.34,
        0.20,
        12,
        MUTED,
        False,
        PP_ALIGN.CENTER,
    )
    output.parent.mkdir(parents=True, exist_ok=True)
    prs.save(output)


def draw_centered(
    draw: ImageDraw.ImageDraw,
    value: str,
    box: tuple[int, int, int, int],
    font: ImageFont.FreeTypeFont,
    fill: tuple[int, int, int],
) -> None:
    x0, y0, x1, y1 = box
    bounds = draw.textbbox((0, 0), value, font=font)
    width = bounds[2] - bounds[0]
    height = bounds[3] - bounds[1]
    draw.text(
        (x0 + (x1 - x0 - width) / 2, y0 + (y1 - y0 - height) / 2 - bounds[1]),
        value,
        font=font,
        fill=fill,
    )


def build_preview(output: Path) -> None:
    canvas = Image.new("RGB", (1600, 900), "white")
    draw = ImageDraw.Draw(canvas)
    blue = (0, 97, 170)
    deep_blue = (0, 69, 122)
    cyan = (0, 143, 174)
    green = (34, 128, 94)
    ink = (25, 40, 55)
    muted = (81, 101, 118)
    line = (205, 215, 223)
    pale_blue = (236, 246, 252)
    pale_cyan = (235, 248, 250)
    pale_green = (237, 247, 242)
    pale_gray = (245, 247, 249)
    red = (181, 45, 55)

    draw.rectangle((0, 0, 382, 58), fill=blue)
    draw.text((42, 10), "玄女月度地理嵌入", font=pil_font(25, True), fill="white")
    draw.rectangle((0, 58, 1600, 139), fill=blue)
    draw_centered(
        draw,
        "从 AEF 思路到城市尺度月度嵌入",
        (0, 58, 1600, 139),
        pil_font(37, True),
        (255, 255, 255),
    )
    draw_centered(
        draw,
        "参考多源压缩重建与 64 维嵌入场范式，针对月度云雾、小区域训练和细粒度目标定向适配",
        (50, 140, 1550, 181),
        pil_font(21, True),
        ink,
    )

    draw.rectangle((58, 190, 1542, 315), fill=pale_blue, outline=line, width=2)
    draw.text(
        (79, 196),
        "共同底座｜参考 AEF 的 embedding field 范式",
        font=pil_font(17, True),
        fill=blue,
    )
    x0, y, width, gap = 84, 235, 318, 46
    for index, (title, detail) in enumerate(BASE_NODES):
        x = x0 + index * (width + gap)
        draw.rectangle((x, y, x + width, y + 60), fill="white", outline=line, width=2)
        draw_centered(draw, title, (x, y + 2, x + width, y + 31), pil_font(16, True), deep_blue)
        draw_centered(draw, detail, (x, y + 29, x + width, y + 57), pil_font(13), muted)
        if index < len(BASE_NODES) - 1:
            draw_centered(
                draw, "→", (x + width, y, x + width + gap, y + 60), pil_font(20, True), blue
            )

    cards = [
        (58, "01", "月度多源嵌入", pale_blue, blue),
        (564, "02", "困难重建与缺失鲁棒", pale_cyan, cyan),
        (1070, "03", "语义稳定与高分细节", pale_green, green),
    ]
    for x, number, title, fill, accent in cards:
        draw.rectangle((x, 340, x + 474, 793), fill=fill, outline=line, width=2)
        draw.rectangle((x, 340, x + 66, 408), fill=accent)
        draw_centered(draw, number, (x, 340, x + 66, 408), pil_font(18, True), (255, 255, 255))
        draw.text((x + 83, 356), title, font=pil_font(22, True), fill=accent)

    draw_centered(
        draw,
        "按月组织多源观测，不依赖年度合成",
        (82, 417, 508, 461),
        pil_font(16, True),
        ink,
    )
    draw.line((116, 512, 476, 512), fill=blue, width=3)
    for index, month in enumerate(MONTHS):
        cx = 120 + index * 70
        draw.rectangle((cx, 496, cx + 34, 530), fill=blue)
        draw_centered(draw, month, (cx - 8, 536, cx + 42, 568), pil_font(13), muted)
    draw.rectangle((109, 602, 481, 670), fill="white", outline=line, width=2)
    draw.text((128, 622), "每期输出", font=pil_font(14, True), fill=muted)
    draw.text((273, 614), "128×128×64", font=pil_font(25, True), fill=blue)
    draw_centered(
        draw,
        "连续 6 期月度索引｜跨月份注意力融合上下文",
        (90, 700, 500, 766),
        pil_font(15, True),
        ink,
    )

    draw_centered(
        draw,
        "先故意拿走一部分输入，再要求模型补回来",
        (590, 413, 1005, 451),
        pil_font(16, True),
        ink,
    )
    draw_centered(
        draw,
        "整源模态遮挡｜S2 18%｜S1 / Landsat 35%",
        (590, 452, 1005, 486),
        pil_font(13, True),
        cyan,
    )
    row_y = 490
    for title, detail in MASK_FACTS:
        draw.rectangle((593, row_y, 607, row_y + 48), fill=cyan)
        draw.text((621, row_y + 8), title, font=pil_font(14, True), fill=cyan)
        draw.multiline_text(
            (748, row_y + 1),
            detail,
            font=pil_font(10, True),
            fill=ink,
            spacing=2,
        )
        row_y += 58
    draw.rectangle((600, 738, 1002, 776), fill="white", outline=line, width=2)
    draw_centered(
        draw,
        "目的：在云雾、缺月或缺源时仍生成稳定嵌入",
        (604, 738, 998, 774),
        pil_font(12, True),
        deep_blue,
    )

    draw_centered(
        draw,
        "OSM 只做语义提示，不把下游真值写进底座",
        (1095, 417, 1519, 461),
        pil_font(16, True),
        ink,
    )
    draw_centered(
        draw,
        "13 项可重叠的 OSM 语义提示\n建筑、主路、支路、铁路、水体、绿地、农业、居住\n"
        "商业、工业、施工、步道、操场",
        (1095, 458, 1519, 526),
        pil_font(12, True),
        ink,
    )
    draw.rectangle((1102, 535, 1510, 619), fill="white", outline=line, width=2)
    draw_centered(
        draw,
        "临时 1×1 线性探针对每个像素分别判断\n是否属于各项可重叠语义；训练结束后丢弃",
        (1108, 540, 1504, 614),
        pil_font(12, True),
        deep_blue,
    )
    draw_centered(
        draw,
        "0.14 控制整组 OSM 提示对训练的影响力度\n"
        "不是准确率、数据占比，也不代表总损失的 14%",
        (1095, 626, 1519, 681),
        pil_font(12, True),
        green,
    )
    draw_centered(
        draw,
        "前 80 轮逐步加到 0.14",
        (1095, 690, 1519, 744),
        pil_font(12, True),
        muted,
    )

    draw.rectangle((58, 811, 1542, 855), fill=pale_gray)
    draw_centered(
        draw,
        "证据边界｜不是完全复现 AEF；OSM 是辅助弱语义，不是人工真值；"
        "P10C 中人工变化监督 0、时间对比损失 0。",
        (70, 813, 1530, 853),
        pil_font(14, True),
        red,
    )
    draw_centered(
        draw,
        "本页参数均来自海淀 P10C epoch 800 配置；不宣称全国或跨区域零样本泛化。",
        (60, 863, 1540, 892),
        pil_font(12),
        muted,
    )
    output.parent.mkdir(parents=True, exist_ok=True)
    canvas.save(output, optimize=True)


def main() -> None:
    args = parse_args()
    build_ppt(args.source_pptx, args.output_pptx)
    build_preview(args.preview)
    metadata = {
        "slide": 3,
        "title": "从 AEF 思路到城市尺度月度嵌入",
        "evidence": {
            "config": (
                "configs/v2_p10c_haidian_202512_202605_osm_semantic_hardneg_" "20260704.yaml"
            ),
            "model": "P10C epoch 800",
            "embed_shape": [128, 128, 64],
            "months": 6,
            "masking": {
                "modality_dropout": {"s2": 0.18, "s1": 0.35, "landsat": 0.35},
                "month_dropout": 0.65,
                "max_months_per_sample": 4,
                "spatial_block_probability": 0.65,
                "spatial_block_ratio": 0.32,
                "drop_availability_masks": False,
                "availability_mask_role": "source-month availability for masking and attention",
                "pixel_mask_role": (
                    "valid pixels for aggregation, reconstruction and high-res fusion"
                ),
            },
            "highres_fusion": "aggregate observations to one frame, then inject into each month",
            "loss_weights": {
                "uniformity": 0.06,
                "osm_semantic_probe": 0.14,
                "osm_hard_negative_ratio": 0.02,
                "osm_hard_negative_weight": 0.35,
                "highres_optical_reconstruction": 0.90,
                "s1_reconstruction": 0.25,
                "highres_sar_reconstruction": 0.35,
                "supervised_change": 0.0,
                "temporal_contrast": 0.0,
            },
        },
        "claim_boundary": [
            "AEF-inspired embedding field paradigm, not a complete AEF reproduction.",
            "OSM is auxiliary weak semantics, not manual ground truth.",
            "No national or cross-region zero-shot generalization claim.",
        ],
    }
    args.asset_root.mkdir(parents=True, exist_ok=True)
    (args.asset_root / "metadata.json").write_text(
        json.dumps(metadata, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    print(args.output_pptx)


if __name__ == "__main__":
    main()
