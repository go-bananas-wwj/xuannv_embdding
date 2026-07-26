"""Shared native PowerPoint primitives for the editable Pujiang deck."""

from __future__ import annotations

from copy import deepcopy
from pathlib import Path
from typing import Literal

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_SHAPE_TYPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.util import Inches, Pt

SLIDE_WIDTH = Inches(13.333333)
SLIDE_HEIGHT = Inches(7.5)

BLUE = RGBColor(0, 97, 170)
DEEP_BLUE = RGBColor(0, 69, 122)
CYAN = RGBColor(0, 143, 174)
GREEN = RGBColor(34, 128, 94)
INK = RGBColor(25, 40, 55)
MUTED = RGBColor(81, 101, 118)
LINE = RGBColor(205, 215, 223)
PALE_BLUE = RGBColor(236, 246, 252)
PALE_CYAN = RGBColor(235, 248, 250)
PALE_GREEN = RGBColor(237, 247, 242)
PALE_GRAY = RGBColor(245, 247, 249)
WHITE = RGBColor(255, 255, 255)
RED = RGBColor(181, 45, 55)
FONT = "Microsoft YaHei"


def new_blank_slide(prs: Presentation):
    """Append a blank 16:9 slide to ``prs`` without placeholders."""
    prs.slide_width = SLIDE_WIDTH
    prs.slide_height = SLIDE_HEIGHT
    return prs.slides.add_slide(prs.slide_layouts[6])


def add_text(
    slide,
    value: str,
    x: float,
    y: float,
    w: float,
    h: float,
    size: float,
    color: RGBColor = INK,
    bold: bool = False,
    align: PP_ALIGN = PP_ALIGN.LEFT,
    font: str = FONT,
    vertical_anchor: MSO_ANCHOR = MSO_ANCHOR.MIDDLE,
):
    """Add an editable, word-wrapped text box using inch-based geometry."""
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.vertical_anchor = vertical_anchor
    frame.margin_left = 0
    frame.margin_right = 0
    frame.margin_top = 0
    frame.margin_bottom = 0

    lines = value.split("\n")
    for index, line in enumerate(lines):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.alignment = align
        paragraph.space_after = Pt(0)
        run = paragraph.add_run()
        run.text = line
        run.font.name = font
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = color
    return shape


def add_rect(
    slide,
    x: float,
    y: float,
    w: float,
    h: float,
    fill: RGBColor,
    line: RGBColor | None = None,
    shape_type: MSO_SHAPE = MSO_SHAPE.RECTANGLE,
):
    """Add an editable native PowerPoint rectangle or rounded rectangle."""
    shape = slide.shapes.add_shape(shape_type, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = line or fill
    return shape


def add_header(slide, title: str, kicker: str):
    """Add the deck's shared editable title band and supporting kicker."""
    add_rect(slide, 0.0, 0.0, 13.333333, 0.82, BLUE)
    add_text(slide, title, 0.55, 0.13, 8.8, 0.36, 28, WHITE, True)
    add_text(slide, kicker, 0.57, 0.51, 11.9, 0.22, 12, PALE_BLUE)


def add_picture(
    slide,
    path: str | Path,
    x: float,
    y: float,
    w: float,
    h: float,
    fit: Literal["contain", "cover"] = "contain",
):
    """Add an independent image object fitted to an inch-based bounding box."""
    image_path = Path(path)
    with Image.open(image_path) as image:
        image_width, image_height = image.size
    if image_width <= 0 or image_height <= 0:
        raise ValueError(f"Image has invalid dimensions: {image_path}")
    if fit not in {"contain", "cover"}:
        raise ValueError(f"Unsupported image fit: {fit}")

    box_ratio = w / h
    image_ratio = image_width / image_height
    if fit == "contain":
        if image_ratio > box_ratio:
            width, height = w, w / image_ratio
        else:
            width, height = h * image_ratio, h
        return slide.shapes.add_picture(
            str(image_path),
            Inches(x + (w - width) / 2),
            Inches(y + (h - height) / 2),
            width=Inches(width),
            height=Inches(height),
        )

    picture = slide.shapes.add_picture(
        str(image_path), Inches(x), Inches(y), width=Inches(w), height=Inches(h)
    )
    if image_ratio > box_ratio:
        crop = (1 - box_ratio / image_ratio) / 2
        picture.crop_left = crop
        picture.crop_right = crop
    elif image_ratio < box_ratio:
        crop = (1 - image_ratio / box_ratio) / 2
        picture.crop_top = crop
        picture.crop_bottom = crop
    return picture


def is_full_slide_picture(shape, prs: Presentation, tolerance: int = 1) -> bool:
    """Return whether ``shape`` is a picture covering the full slide geometry."""
    return (
        shape.shape_type == MSO_SHAPE_TYPE.PICTURE
        and abs(shape.left) <= tolerance
        and abs(shape.top) <= tolerance
        and abs(shape.width - prs.slide_width) <= tolerance
        and abs(shape.height - prs.slide_height) <= tolerance
    )


def clone_first_slide(source: str | Path | Presentation, destination: Presentation):
    """Clone the first source slide into ``destination``, including picture relations."""
    source_presentation = Presentation(source) if isinstance(source, (str, Path)) else source
    if not source_presentation.slides:
        raise ValueError("Source presentation has no slides")

    source_slide = source_presentation.slides[0]
    destination.slide_width = source_presentation.slide_width
    destination.slide_height = source_presentation.slide_height
    cloned_slide = new_blank_slide(destination)

    relationship_ids: dict[str, str] = {}
    for relationship in source_slide.part.rels.values():
        if relationship.reltype == RT.SLIDE_LAYOUT:
            continue
        target = relationship.target_ref if relationship.is_external else relationship.target_part
        relationship_ids[relationship.rId] = cloned_slide.part.relate_to(
            target, relationship.reltype, relationship.is_external
        )

    for shape in source_slide.shapes:
        cloned_element = deepcopy(shape.element)
        for element in cloned_element.iter():
            for attribute, value in tuple(element.attrib.items()):
                if value in relationship_ids:
                    element.set(attribute, relationship_ids[value])
        cloned_slide.shapes._spTree.insert_element_before(cloned_element, "p:extLst")
    return cloned_slide
