#!/usr/bin/env python3
"""Build Pujiang presentation slide 2: monthly sample preparation and training facts."""

from __future__ import annotations

import argparse
import json
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

DEFAULT_SOURCE_PPTX = Path(
    "/root/.codex/attachments/d0a9024d-03de-463f-83fa-675e2a235467/" "Alpha Earth介绍.pptx"
)
DEFAULT_QA_IMAGE = Path(
    "/data/xuannv_embedding/qa/pretrain_p6a_20260703/"
    "training_seen_batches_clean/patch_000183_clean_training_tensor_grid.png"
)
DEFAULT_ASSET_ROOT = Path("/root/workspace/xuannv/docs/presentations/assets/pujiang_202607/slide02")
DEFAULT_OUTPUT = Path(
    "/root/workspace/xuannv/docs/presentations/pujiang_202607/"
    "玄女月度地理嵌入_浦江交流_第02页_20260726.pptx"
)
DEFAULT_PREVIEW = DEFAULT_ASSET_ROOT / "slide02_preview.png"

BLUE = RGBColor(0, 97, 170)
DEEP_BLUE = RGBColor(0, 69, 122)
INK = RGBColor(25, 40, 55)
MUTED = RGBColor(81, 101, 118)
LINE = RGBColor(205, 215, 223)
PALE_BLUE = RGBColor(236, 246, 252)
PALE_GRAY = RGBColor(245, 247, 249)
WHITE = RGBColor(255, 255, 255)
RED = RGBColor(181, 45, 55)
FONT = "Microsoft YaHei"

MONTHS = ["2025.12", "2026.01", "2026.02", "2026.03", "2026.04", "2026.05"]
SOURCE_ROWS = [
    ("Sentinel-2", (183, 459)),
    ("高分光学", (1079, 1353)),
    ("高分SAR", (1376, 1651)),
]
SOURCE_COLUMNS = [
    (185, 460),
    (675, 950),
    (1170, 1445),
    (1665, 1940),
    (2160, 2435),
    (2655, 2930),
]
SOURCE_SUMMARY = "训练数据｜S2、S1、Landsat、高分光学、高分SAR、清洗后的13类 OSM 弱语义"
PROCESS_STEPS = [
    "区域网格划分",
    "多源影像检索",
    "云雾与无效像素筛选",
    "空间配准",
    "月度有效像素聚合",
    "训练样本 QA",
]
HAIDIAN_FACTS = [
    ("空间样本", "320 patch｜1280 m × 1280 m"),
    ("时间范围", "2025.12–2026.05｜连续6期"),
    ("训练输入", "六类多源数据＋像素级 mask"),
    ("训练设置", "P10C epoch 800｜2×Ascend NPU"),
    ("实测时长", "约 17 小时 50 分"),
    ("初始化", "从 P9B epoch 800 继续训练"),
]
HARBIN_FACTS = [
    ("空间样本", "424 patch"),
    ("API可用月份", "10期｜2025.04–2026.05（非连续）"),
    ("代表模型", "V5（API v2）"),
    ("训练数据", "API 未披露"),
    ("训练硬件", "API 未披露"),
    ("训练时长", "API 未披露"),
]


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source-pptx", type=Path, default=DEFAULT_SOURCE_PPTX)
    parser.add_argument("--qa-image", type=Path, default=DEFAULT_QA_IMAGE)
    parser.add_argument("--asset-root", type=Path, default=DEFAULT_ASSET_ROOT)
    parser.add_argument("--output-pptx", type=Path, default=DEFAULT_OUTPUT)
    parser.add_argument("--preview", type=Path, default=DEFAULT_PREVIEW)
    return parser.parse_args()


def pil_font(size: int, bold: bool = False) -> ImageFont.FreeTypeFont:
    candidates = [
        (
            "/usr/share/fonts/google-noto-cjk/NotoSansCJK-Bold.ttc"
            if bold
            else "/usr/share/fonts/google-noto-cjk/NotoSansCJK-Regular.ttc"
        ),
        (
            "/usr/share/fonts/truetype/wqy/wqy-zenhei.ttc"
            if not bold
            else "/usr/share/fonts/truetype/wqy/wqy-zenhei.ttc"
        ),
        "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
    ]
    for path in candidates:
        if Path(path).exists():
            return ImageFont.truetype(path, size)
    return ImageFont.load_default()


def build_monthly_input_montage(source: Path, output: Path) -> dict[str, list[str] | str]:
    """Create a compact 3-source × 6-month view from the audited training-input sheet."""
    tile_w, tile_h = 252, 166
    left, top = 210, 76
    gap_x, gap_y = 12, 18
    canvas = Image.new("RGB", (1800, 630), "white")
    draw = ImageDraw.Draw(canvas)
    title_font = pil_font(26, True)
    label_font = pil_font(21, True)
    small_font = pil_font(19)
    draw.text(
        (22, 18),
        "真实训练输入示例｜patch 000183",
        fill=(25, 40, 55),
        font=title_font,
    )
    with Image.open(source) as image:
        image = image.convert("RGB")
        for col, (month, (x0, x1)) in enumerate(zip(MONTHS, SOURCE_COLUMNS, strict=True)):
            x = left + col * (tile_w + gap_x)
            month_box = draw.textbbox((0, 0), month, font=small_font)
            draw.text(
                (x + (tile_w - month_box[2] + month_box[0]) / 2, 49),
                month,
                fill=(81, 101, 118),
                font=small_font,
            )
            for row, (source_name, (y0, y1)) in enumerate(SOURCE_ROWS):
                y = top + row * (tile_h + gap_y)
                crop = image.crop((x0, y0, x1, y1))
                crop = crop.resize((tile_w, tile_h), Image.Resampling.LANCZOS)
                canvas.paste(crop, (x, y))
                draw.rectangle((x, y, x + tile_w, y + tile_h), outline=(205, 215, 223), width=2)
                if col == 0:
                    label_box = draw.textbbox((0, 0), source_name, font=label_font)
                    draw.text(
                        (
                            left - 18 - (label_box[2] - label_box[0]),
                            y + (tile_h - label_box[3] + label_box[1]) / 2,
                        ),
                        source_name,
                        fill=(0, 97, 170),
                        font=label_font,
                    )
    output.parent.mkdir(parents=True, exist_ok=True)
    canvas.save(output, optimize=True)
    return {
        "source": str(source),
        "months": MONTHS,
        "sources": [item[0] for item in SOURCE_ROWS],
    }


def clear_slide(slide) -> None:
    for shape in list(slide.shapes):
        shape._element.getparent().remove(shape._element)


def keep_only_slide(prs: Presentation, index: int) -> None:
    slide_ids = prs.slides._sldIdLst
    keep = slide_ids[index]
    for slide_id in list(slide_ids):
        if slide_id is not keep:
            slide_ids.remove(slide_id)


def add_rect(slide, x: float, y: float, w: float, h: float, fill, line=None):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(x),
        Inches(y),
        Inches(w),
        Inches(h),
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = line or fill
    return shape


def add_text(
    slide,
    value: str,
    x: float,
    y: float,
    w: float,
    h: float,
    size: int,
    color=INK,
    bold: bool = False,
    align=PP_ALIGN.LEFT,
):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    paragraph.space_after = Pt(0)
    run = paragraph.add_run()
    run.text = value
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return shape


def add_picture_contain(slide, path: Path, x: float, y: float, w: float, h: float):
    with Image.open(path) as image:
        ratio = image.width / image.height
    box_ratio = w / h
    if ratio >= box_ratio:
        draw_w, draw_h = w, w / ratio
    else:
        draw_h, draw_w = h, h * ratio
    return slide.shapes.add_picture(
        str(path),
        Inches(x + (w - draw_w) / 2),
        Inches(y + (h - draw_h) / 2),
        Inches(draw_w),
        Inches(draw_h),
    )


def add_source_badges(slide) -> None:
    badges = [
        ("S2", "多光谱光学"),
        ("S1", "雷达结构"),
        ("Landsat", "跨传感器光学"),
        ("高分光学", "边界与小目标"),
        ("高分SAR", "精细散射结构"),
        ("OSM", "13类弱语义"),
    ]
    x0, y, width, gap = 0.48, 1.58, 1.31, 0.08
    for index, (name, role) in enumerate(badges):
        x = x0 + index * (width + gap)
        add_rect(slide, x, y, width, 0.52, PALE_BLUE, LINE)
        add_text(
            slide, name, x + 0.05, y + 0.02, width - 0.10, 0.22, 12, BLUE, True, PP_ALIGN.CENTER
        )
        add_text(
            slide, role, x + 0.04, y + 0.25, width - 0.08, 0.20, 8, MUTED, False, PP_ALIGN.CENTER
        )


def add_process(slide) -> None:
    x0, y, width, gap = 0.48, 2.57, 1.28, 0.11
    for index, label in enumerate(PROCESS_STEPS):
        number = f"{index + 1:02d}"
        x = x0 + index * (width + gap)
        add_rect(slide, x, y, width, 0.73, WHITE, LINE)
        add_rect(slide, x, y, 0.30, 0.73, BLUE, BLUE)
        add_text(slide, number, x, y + 0.04, 0.30, 0.65, 11, WHITE, True, PP_ALIGN.CENTER)
        add_text(
            slide, label, x + 0.35, y + 0.05, width - 0.40, 0.63, 11, INK, True, PP_ALIGN.CENTER
        )
        if index < len(PROCESS_STEPS) - 1:
            add_text(slide, "→", x + width, y + 0.17, gap, 0.38, 14, BLUE, True, PP_ALIGN.CENTER)


def add_fact_row(slide, label: str, value: str, x: float, y: float, value_color=INK) -> None:
    add_text(slide, label, x, y, 1.12, 0.31, 11, MUTED)
    add_text(slide, value, x + 1.06, y, 2.83, 0.31, 12, value_color, True)


def add_region_facts(
    slide,
    y: float,
    title: str,
    rows: list[tuple[str, str]],
    accent=BLUE,
    spacing: float = 0.34,
) -> None:
    add_rect(slide, 8.86, y, 4.00, 0.40, accent, accent)
    add_text(slide, title, 8.99, y + 0.01, 3.74, 0.37, 15, WHITE, True)
    row_y = y + 0.49
    for label, value in rows:
        add_fact_row(slide, label, value, 8.99, row_y)
        row_y += spacing


def draw_centered(
    draw: ImageDraw.ImageDraw,
    value: str,
    box: tuple[int, int, int, int],
    font: ImageFont.FreeTypeFont,
    fill: tuple[int, int, int],
) -> None:
    x0, y0, x1, y1 = box
    bounds = draw.textbbox((0, 0), value, font=font)
    width = bounds[2] - bounds[0]
    height = bounds[3] - bounds[1]
    draw.text(
        (x0 + (x1 - x0 - width) / 2, y0 + (y1 - y0 - height) / 2 - bounds[1]),
        value,
        font=font,
        fill=fill,
    )


def paste_contain(canvas: Image.Image, source: Path, box: tuple[int, int, int, int]) -> None:
    x0, y0, x1, y1 = box
    with Image.open(source) as image:
        image = image.convert("RGB")
        image.thumbnail((x1 - x0, y1 - y0), Image.Resampling.LANCZOS)
        x = x0 + (x1 - x0 - image.width) // 2
        y = y0 + (y1 - y0 - image.height) // 2
        canvas.paste(image, (x, y))


def build_preview(output: Path, montage: Path) -> None:
    """Render a deterministic review image using the same slide coordinates."""
    canvas = Image.new("RGB", (1600, 900), "white")
    draw = ImageDraw.Draw(canvas)
    blue = (0, 97, 170)
    deep_blue = (0, 69, 122)
    ink = (25, 40, 55)
    muted = (81, 101, 118)
    line = (205, 215, 223)
    pale_blue = (236, 246, 252)
    pale_gray = (245, 247, 249)
    red = (181, 45, 55)

    draw.rectangle((0, 0, 382, 58), fill=blue)
    draw.text((42, 10), "玄女月度地理嵌入", font=pil_font(25, True), fill="white")
    draw.rectangle((0, 58, 1600, 139), fill=blue)
    draw_centered(
        draw,
        "两个区域的多源数据处理与训练组织",
        (0, 58, 1600, 139),
        pil_font(37, True),
        (255, 255, 255),
    )
    draw_centered(
        draw,
        "海淀 P10C 完整追溯数据与训练流程；哈尔滨 V5 仅采用 API 已披露的部署信息",
        (50, 140, 1550, 181),
        pil_font(22, True),
        ink,
    )

    draw.text((58, 190), "左侧流程与样例：海淀 P10C", font=pil_font(18, True), fill=blue)
    draw.rectangle((58, 229, 1032, 281), fill=pale_blue, outline=line, width=2)
    draw_centered(draw, SOURCE_SUMMARY, (68, 232, 1022, 278), pil_font(14, True), deep_blue)

    sx0, sy, sw, sgap = 58, 298, 154, 13
    for index, label in enumerate(PROCESS_STEPS):
        x = sx0 + index * (sw + sgap)
        draw.rectangle((x, sy, x + sw, sy + 82), fill="white", outline=line, width=2)
        draw.rectangle((x, sy, x + 36, sy + 82), fill=blue)
        draw_centered(
            draw,
            f"{index + 1:02d}",
            (x, sy, x + 36, sy + 82),
            pil_font(14, True),
            (255, 255, 255),
        )
        draw_centered(
            draw,
            label,
            (x + 40, sy + 4, x + sw - 4, sy + 78),
            pil_font(13, True),
            ink,
        )
        if index < len(PROCESS_STEPS) - 1:
            draw_centered(
                draw,
                "→",
                (x + sw, sy + 20, x + sw + sgap, sy + 68),
                pil_font(16, True),
                blue,
            )

    draw.rectangle((58, 393, 1032, 809), fill="white", outline=line, width=2)
    paste_contain(canvas, montage, (70, 402, 1020, 715))
    draw.rectangle((82, 730, 1008, 790), fill=pale_gray)
    draw_centered(
        draw,
        "像素级 mask 排除云、阴影和缺测；同月多景仅对有效像素做加权聚合",
        (90, 735, 1000, 787),
        pil_font(15, True),
        deep_blue,
    )

    def region_block(
        y: int,
        title: str,
        rows: list[tuple[str, str]],
        color: tuple[int, int, int],
        spacing: int = 40,
    ) -> None:
        draw.rectangle((1063, y, 1544, y + 48), fill=color)
        draw.text((1078, y + 10), title, font=pil_font(18, True), fill="white")
        row_y = y + 59
        for label, value in rows:
            draw.text((1078, row_y + 2), label, font=pil_font(12), fill=muted)
            draw.text((1205, row_y + 2), value, font=pil_font(12, True), fill=ink)
            row_y += spacing

    region_block(
        190,
        "海淀区｜P10C 生产版",
        HAIDIAN_FACTS,
        blue,
    )
    region_block(
        497,
        "哈尔滨新区｜API 已部署产品",
        HARBIN_FACTS,
        deep_blue,
    )
    draw.rectangle((1063, 810, 1544, 843), fill=pale_gray)
    draw_centered(
        draw,
        "V5 训练记录未接入本页；未混用旧 Stage 2 训练记录",
        (1070, 811, 1537, 841),
        pil_font(11, True),
        red,
    )
    draw_centered(
        draw,
        "海淀 P10C 质量筛选：S2 使用 SCL，Landsat 使用 QA_PIXEL；16,339 景候选中保留 6,077 景。",
        (45, 851, 1555, 888),
        pil_font(14),
        muted,
    )
    output.parent.mkdir(parents=True, exist_ok=True)
    canvas.save(output, optimize=True)


def build_ppt(source: Path, output: Path, montage: Path) -> None:
    prs = Presentation(source)
    keep_only_slide(prs, 1)
    slide = prs.slides[0]
    clear_slide(slide)
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = WHITE

    add_rect(slide, 0.0, 0.0, 3.18, 0.48, BLUE)
    add_text(slide, "玄女月度地理嵌入", 0.35, 0.02, 2.55, 0.42, 17, WHITE, True)
    add_rect(slide, 0.0, 0.48, 13.333, 0.68, BLUE)
    add_text(
        slide,
        "两个区域的多源数据处理与训练组织",
        0.42,
        0.50,
        12.5,
        0.62,
        29,
        WHITE,
        True,
        PP_ALIGN.CENTER,
    )
    add_text(
        slide,
        "海淀 P10C 完整追溯数据与训练流程；哈尔滨 V5 仅采用 API 已披露的部署信息",
        0.48,
        1.17,
        12.37,
        0.32,
        18,
        INK,
        True,
        PP_ALIGN.CENTER,
    )

    add_text(
        slide,
        "左侧流程与样例：海淀 P10C",
        0.48,
        1.57,
        8.12,
        0.30,
        15,
        BLUE,
        True,
    )
    add_rect(slide, 0.48, 1.94, 8.12, 0.46, PALE_BLUE, LINE)
    add_text(
        slide,
        SOURCE_SUMMARY,
        0.63,
        1.98,
        7.82,
        0.36,
        11,
        DEEP_BLUE,
        True,
        PP_ALIGN.CENTER,
    )
    add_process(slide)
    add_rect(slide, 0.48, 3.44, 8.12, 3.30, WHITE, LINE)
    add_picture_contain(slide, montage, 0.58, 3.52, 7.92, 2.43)
    add_rect(slide, 0.68, 6.10, 7.72, 0.48, PALE_GRAY, PALE_GRAY)
    add_text(
        slide,
        "像素级 mask 排除云、阴影和缺测；同月多景仅对有效像素做加权聚合",
        0.83,
        6.13,
        7.42,
        0.40,
        12,
        DEEP_BLUE,
        True,
        PP_ALIGN.CENTER,
    )

    add_region_facts(
        slide,
        1.58,
        "海淀区｜P10C 生产版",
        HAIDIAN_FACTS,
    )
    add_region_facts(
        slide,
        4.16,
        "哈尔滨新区｜API 已部署产品",
        HARBIN_FACTS,
        DEEP_BLUE,
        0.30,
    )
    add_rect(slide, 8.86, 6.54, 4.00, 0.37, PALE_GRAY, PALE_GRAY)
    add_text(
        slide,
        "V5 训练记录未接入本页；未混用旧 Stage 2 训练记录",
        9.00,
        6.55,
        3.72,
        0.33,
        10,
        RED,
        True,
        PP_ALIGN.CENTER,
    )
    add_text(
        slide,
        "海淀 P10C 质量筛选：S2 使用 SCL，Landsat 使用 QA_PIXEL；16,339 景候选中保留 6,077 景。",
        0.50,
        7.02,
        12.34,
        0.25,
        11,
        MUTED,
        False,
        PP_ALIGN.CENTER,
    )
    output.parent.mkdir(parents=True, exist_ok=True)
    prs.save(output)


def main() -> None:
    args = parse_args()
    montage = args.asset_root / "haidian_patch_000183_monthly_training_inputs.png"
    montage_metadata = build_monthly_input_montage(args.qa_image, montage)
    build_ppt(args.source_pptx, args.output_pptx, montage)
    build_preview(args.preview, montage)
    metadata = {
        "slide": 2,
        "title": "两个区域的多源数据处理与训练组织",
        "montage": montage_metadata,
        "haidian": {
            "patches": 320,
            "months": MONTHS,
            "model": "P10C epoch 800",
            "initialization": "P9B epoch 800",
            "hardware": "2 Ascend NPU",
            "elapsed": "17:49:31",
            "training_log": "/data/xuannv_embedding/logs/p10c_train_20260704.log",
        },
        "harbin": {
            "patches": 424,
            "available_month_count": 10,
            "available_months_are_training_period": False,
            "model": "V5 (API v2)",
            "training_data": "API not disclosed",
            "hardware": "API not disclosed",
            "elapsed": "API not disclosed",
            "evidence_boundary": "API deployment facts only; no legacy Stage 2 training claims.",
        },
    }
    args.asset_root.mkdir(parents=True, exist_ok=True)
    (args.asset_root / "metadata.json").write_text(
        json.dumps(metadata, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    print(args.output_pptx)


if __name__ == "__main__":
    main()
